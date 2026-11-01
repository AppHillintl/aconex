import io
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import duckdb

from pathlib import Path as _Path
_SRV_ROOT = _Path(__file__).resolve().parent.parent.parent
_SRV_CLOUD_DB = _SRV_ROOT / "aconex_cloud.duckdb"
_SRV_LOCAL_DB = _SRV_ROOT / "aconex.duckdb"

if _SRV_LOCAL_DB.exists():
    DB_PATH = str(_SRV_LOCAL_DB).replace("\\", "/")
elif _SRV_CLOUD_DB.exists():
    DB_PATH = str(_SRV_CLOUD_DB).replace("\\", "/")
else:
    DB_PATH = str(_SRV_LOCAL_DB).replace("\\", "/")

def get_db_connection():
    return duckdb.connect(DB_PATH, read_only=True)

def generate_filtered_excel_export(
    project_id: str,
    category: str,
    search: str = None,
    status: str = None,
    discipline: str = None,
    doc_type: str = None,
    mail_type: str = None,
    is_overdue: str = None,
    sort: str = None,
    from_org: str = None,
    to_org: str = None,
    direction: str = None,
    limit: int = 50000
) -> tuple[bytes, str]:
    """
    Generates a professionally styled Excel workbook containing exactly the filtered and sorted data
    for the selected category (documents, mail, rfis, workflows).
    """
    con = get_db_connection()
    try:
        # 1. Fetch Project Metadata
        proj_info = con.execute("""
            SELECT project_id, project_title, k_package_no 
            FROM dim_project_packages 
            WHERE project_id = ? 
            LIMIT 1
        """, [project_id]).fetchone()
        
        proj_name = proj_info[1] if proj_info else project_id
        pkg_code = proj_info[2] if proj_info and proj_info[2] else project_id

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.views.sheetView[0].showGridLines = True

        cat_clean = (category or 'documents').lower()

        # Styles
        title_font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
        title_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        sub_font = Font(name="Calibri", size=10, italic=True, color="94A3B8")
        sub_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        
        link_font = Font(name="Calibri", size=10, bold=True, color="0284C7", underline="single")
        regular_font = Font(name="Calibri", size=10, color="0F172A")
        bold_font = Font(name="Calibri", size=10, bold=True, color="0F172A")
        
        thin_border_side = Side(style='thin', color='E2E8F0')
        thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
        
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        
        overdue_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        overdue_font = Font(name="Calibri", size=10, bold=True, color="991B1B")
        
        closed_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
        closed_font = Font(name="Calibri", size=10, bold=True, color="166534")

        # 2. Query data based on category
        filter_summary_parts = []
        if search: filter_summary_parts.append(f"Search: '{search}'")
        if direction: filter_summary_parts.append(f"Direction: {'From Contractor (Inbound)' if direction == 'inbound' else ('To Contractor (Outbound)' if direction == 'outbound' else direction)}")
        if from_org: filter_summary_parts.append(f"From Org: {from_org}")
        if to_org: filter_summary_parts.append(f"To Org: {to_org}")
        if status: filter_summary_parts.append(f"Status: {status}")
        if discipline: filter_summary_parts.append(f"Discipline: {discipline}")
        if doc_type: filter_summary_parts.append(f"Doc Type: {doc_type}")
        if mail_type: filter_summary_parts.append(f"Mail Type: {mail_type}")
        if is_overdue in ['true', 'yes', '1']: filter_summary_parts.append("Overdue Only: Yes")
        if sort: filter_summary_parts.append(f"Sorted By: {sort}")

        filter_desc = " | ".join(filter_summary_parts) if filter_summary_parts else "All Records (No Filters Active)"

        if cat_clean in ['documents', 'shop_drawings', 'material_submittals', 'other_docs']:
            if cat_clean == 'shop_drawings':
                ws.title = "Shop Drawings"
            elif cat_clean == 'material_submittals':
                ws.title = "Material Submittals"
            elif cat_clean == 'other_docs':
                ws.title = "Other Documents"
            else:
                ws.title = "Documents Register"
            headers = ["Document No", "Title", "Discipline", "Doc Type", "Status", "Review Status", "Revision", "Version", "Date Modified", "Aconex Link"]
            
            filters = [f"project_package = '{project_id}'", "from_organization NOT ILIKE '%Hill Inter%'"]
            if cat_clean == 'shop_drawings':
                filters.append("(doc_type ILIKE '%Drawing%' OR doc_type IN ('Drawing', 'Shop Drawing', 'As-Built', 'Sketch'))")
            elif cat_clean == 'material_submittals':
                filters.append("(doc_type ILIKE '%Material%' OR title ILIKE '%MAR-%' OR title ILIKE '%MAT-%' OR doc_type ILIKE '%MAR%' OR doc_type ILIKE '%MAT%')")
            elif cat_clean == 'other_docs':
                filters.append("NOT (doc_type ILIKE '%Drawing%' OR doc_type IN ('Drawing', 'Shop Drawing', 'As-Built', 'Sketch')) AND NOT (doc_type ILIKE '%Material%' OR title ILIKE '%MAR-%' OR title ILIKE '%MAT-%' OR doc_type ILIKE '%MAR%' OR doc_type ILIKE '%MAT%')")

            if search:
                s_c = search.replace("'", "''")
                filters.append(f"(document_no ILIKE '%{s_c}%' OR title ILIKE '%{s_c}%' OR discipline_name ILIKE '%{s_c}%')")
            if discipline:
                filters.append(f"discipline_name = '{discipline.replace("'", "''")}'")
            if doc_type:
                filters.append(f"doc_type = '{doc_type.replace("'", "''")}'")
            if status:
                filters.append(f"status = '{status.replace("'", "''")}'")

            order_by = "date_modified DESC NULLS LAST"
            if sort:
                cols_map = {'document_no': 'document_no', 'title': 'title', 'discipline': 'discipline_name', 'type': 'doc_type', 'doc_type': 'doc_type', 'status': 'status', 'revision': 'revision', 'date_modified': 'date_modified'}
                clauses = []
                for t in sort.split(','):
                    parts = t.strip().split()
                    if parts and parts[0].lower() in cols_map:
                        d = 'DESC' if len(parts) > 1 and parts[1].upper() == 'DESC' else 'ASC'
                        clauses.append(f"{cols_map[parts[0].lower()]} {d} NULLS LAST")
                if clauses: order_by = ", ".join(clauses)

            where_sql = f"WHERE {' AND '.join(filters)}"
            rows = con.execute(f"""
                SELECT document_no, title, discipline_name, doc_type, status, review_status, revision, version_number,
                       date_modified::VARCHAR, 
                       CASE 
                           WHEN document_id IS NOT NULL AND LENGTH(TRIM(document_id)) > 3 AND document_id != 'None'
                           THEN 'https://mea.aconex.com/ViewDoc?trackingid=' || document_id || '&projectid=1744833878&docno=' || document_no || '&username=aabdelrahim4&action=ViewDocumentSummary'
                           ELSE 'https://mea.aconex.com/authentication/index.html?Action=View&MENU_TARGET=DOC-SEARCH&ENTITY=PROJECT&ID=1744833878&searchQuery=documentNumber%3A%22' || document_no || '%22'
                       END as aconex_url
                FROM dim_documents
                {where_sql}
                ORDER BY {order_by}
                LIMIT {limit}
            """).fetchall()

        elif cat_clean in ['mail', 'letters', 'other_mail', 'hill_ssc']:
            dir_clean = (direction or '').strip().lower()
            if cat_clean == 'letters':
                if dir_clean == 'outbound':
                    ws.title = "Letters - To Contractor"
                elif dir_clean == 'all':
                    ws.title = "Letters - All Directions"
                else:
                    ws.title = "Letters - From Contractor"
            elif cat_clean == 'hill_ssc':
                ws.title = "Hill & SSC Letters"
            elif cat_clean == 'other_mail':
                ws.title = "Other Correspondence"
            else:
                ws.title = "Correspondence Log"
            headers = ["Mail No", "Direction", "Mail Type", "Subject", "Reference No", "From Org", "To Org", "Sent Date", "Due Date", "Days Late", "Status", "Response Mail No", "Response Date", "Overdue", "Aconex Link"]
            
            filters = []
            if cat_clean == 'hill_ssc':
                filters.append("m.mail_type = 'Letter'")
                filters.append("""(
                    m.from_name ILIKE '%ahmed abdelrahim%' 
                    OR m.from_organization ILIKE '%Shaker%' 
                    OR m.from_organization ILIKE '%ECG%' 
                    OR m.to_organization ILIKE '%Shaker%' 
                    OR m.to_organization ILIKE '%ECG%' 
                    OR m.subject ILIKE '%3120%'
                )""")
                if project_id != 'ALL':
                    filters.append(f"(m.project_package = '{project_id}' OR m.subject ILIKE '%{project_id}%')")
                if dir_clean == 'inbound':
                    filters.append("(m.from_name ILIKE '%ahmed abdelrahim%' OR m.from_organization ILIKE '%Shaker%' OR m.from_organization ILIKE '%ECG%')")
                elif dir_clean == 'outbound':
                    filters.append("(m.subject ILIKE '%3120%' OR m.to_organization ILIKE '%Shaker%' OR m.to_organization ILIKE '%ECG%')")
            else:
                filters.append(f"m.project_package = '{project_id}'")
                if cat_clean == 'letters':
                    filters.append("m.mail_type = 'Letter'")
                    if dir_clean == 'outbound':
                        if project_id == 'DP05':
                            filters.append("(m.to_organization ILIKE '%Orascom%' OR m.from_organization ILIKE '%Hill%' OR m.is_outbound_from_hill = true)")
                        elif project_id == 'DP06':
                            filters.append("(m.to_organization ILIKE '%Hassan Allam%' OR m.from_organization ILIKE '%Hill%' OR m.is_outbound_from_hill = true)")
                        elif project_id == 'DP07':
                            filters.append("(m.to_organization ILIKE '%CSCEC%' OR m.from_organization ILIKE '%Hill%' OR m.is_outbound_from_hill = true)")
                        else:
                            filters.append("(m.is_outbound_from_hill = true OR m.from_organization ILIKE '%Hill%')")
                    elif dir_clean == 'all':
                        pass
                    elif not from_org and not to_org:
                        if project_id == 'DP05':
                            filters.append("(m.from_organization ILIKE '%Orascom%' AND m.from_organization NOT ILIKE '%Hill%')")
                        elif project_id == 'DP06':
                            filters.append("(m.from_organization ILIKE '%Hassan Allam%' AND m.from_organization NOT ILIKE '%Hill%')")
                        elif project_id == 'DP07':
                            filters.append("(m.from_organization ILIKE '%CSCEC%' AND m.from_organization NOT ILIKE '%Hill%')")
                        else:
                            filters.append("m.from_organization NOT ILIKE '%Hill%'")
                elif cat_clean == 'other_mail':
                    filters.append("(m.mail_type != 'Letter' AND m.mail_type NOT ILIKE '%RFI%' AND m.mail_type != 'Request For Information' AND m.mail_type NOT ILIKE '%Clarification%')")

            if search:
                s_c = search.replace("'", "''")
                filters.append(f"(m.mail_no ILIKE '%{s_c}%' OR m.subject ILIKE '%{s_c}%' OR m.from_organization ILIKE '%{s_c}%' OR m.to_organization ILIKE '%{s_c}%' OR m.reference_number ILIKE '%{s_c}%')")
            if from_org:
                filters.append(f"m.from_organization ILIKE '%{from_org.replace("'", "''")}%'")
            if to_org:
                filters.append(f"m.to_organization ILIKE '%{to_org.replace("'", "''")}%'")
            if mail_type:
                filters.append(f"m.mail_type = '{mail_type.replace("'", "''")}'")
            if status:
                filters.append(f"m.status = '{status.replace("'", "''")}'")
            if is_overdue in ['true', 'yes', '1']:
                filters.append("m.is_overdue = true")

            order_by = "m.is_overdue DESC, m.sent_date DESC NULLS LAST"
            if sort:
                cols_map = {'mail_no': 'm.mail_no', 'type': 'm.mail_type', 'mail_type': 'm.mail_type', 'subject': 'm.subject', 'ref_no': 'm.reference_number', 'reference_number': 'm.reference_number', 'from_org': 'm.from_organization', 'to_org': 'm.to_organization', 'sent_date': 'm.sent_date', 'due_date': 'm.due_date', 'status': 'm.status'}
                clauses = []
                for t in sort.split(','):
                    parts = t.strip().split()
                    if parts and parts[0].lower() in cols_map:
                        d = 'DESC' if len(parts) > 1 and parts[1].upper() == 'DESC' else 'ASC'
                        clauses.append(f"{cols_map[parts[0].lower()]} {d} NULLS LAST")
                if clauses: order_by = ", ".join(clauses)

            where_sql = f"WHERE {' AND '.join(filters)}"
            rows = con.execute(f"""
                SELECT m.mail_no,
                       CASE WHEN m.from_organization ILIKE '%Hill%' OR m.is_outbound_from_hill = true THEN 'To Contractor (Outbound)' ELSE 'From Contractor (Inbound)' END AS direction_lbl,
                       m.mail_type, m.subject, COALESCE(m.reference_number, m.mail_no), m.from_organization, m.to_organization,
                       m.sent_date::VARCHAR, m.due_date::VARCHAR, m.days_late, m.status,
                       COALESCE(resp.response_mail_no, '-'),
                       COALESCE(resp.response_sent_date::VARCHAR, '-'),
                       CASE WHEN m.is_overdue THEN 'OVERDUE' ELSE 'ON TRACK' END,
                       COALESCE(
                           CASE WHEN m.aconex_url LIKE '%MiniURL%' THEN m.aconex_url END,
                           CASE WHEN l.aconex_url LIKE '%MiniURL%' THEN l.aconex_url END,
                           'https://mea.aconex.com/authentication/index.html?Action=View&MENU_TARGET=MAIL-SEARCH&ENTITY=PROJECT&ID=1744833878&searchQuery=mailNumber%3A%22' || m.mail_no || '%22'
                       ) as aconex_url
                FROM fact_mail m
                LEFT JOIN mail_aconex_links l ON m.mail_no = l.mail_no
                LEFT JOIN (
                    SELECT mail_no, MAX(response_mail_no) as response_mail_no, MAX(response_sent_date) as response_sent_date
                    FROM fact_mail_responses
                    GROUP BY mail_no
                ) resp ON m.mail_no = resp.mail_no
                {where_sql}
                ORDER BY {order_by}
                LIMIT {limit}
            """).fetchall()

        elif cat_clean == 'rfis':
            if project_id == 'DP01':
                ws.title = "Site Clarifications Log"
                headers = ["Clarification No", "Type", "Subject", "From Org", "To Org", "Status", "Sent Date", "Due Date", "Days Late", "Overdue", "Aconex Link"]
                filters = [
                    "(project_package = 'DP01' OR subject ILIKE '%MD01%')",
                    "mail_type IN ('Site Clarification Request', 'Response to Clarification')"
                ]
                if search:
                    s_c = search.replace("'", "''")
                    filters.append(f"(mail_no ILIKE '%{s_c}%' OR subject ILIKE '%{s_c}%' OR from_organization ILIKE '%{s_c}%')")
                if from_org:
                    filters.append(f"from_organization = '{from_org.replace("'", "''")}'")
                if to_org:
                    filters.append(f"to_organization = '{to_org.replace("'", "''")}'")
                if mail_type:
                    filters.append(f"mail_type = '{mail_type.replace("'", "''")}'")
                if status:
                    filters.append(f"status = '{status.replace("'", "''")}'")
                if is_overdue in ['true', 'yes', '1']:
                    filters.append("is_overdue = true")

                order_by = "is_overdue DESC, sent_date DESC NULLS LAST"
                where_sql = f"WHERE {' AND '.join(filters)}"
                rows = con.execute(f"""
                    SELECT m.mail_no, m.mail_type, m.subject, m.from_organization, m.to_organization,
                           m.status, m.sent_date::VARCHAR, m.due_date::VARCHAR, m.days_late,
                           CASE WHEN m.is_overdue THEN 'OVERDUE' ELSE 'ON TRACK' END,
                           COALESCE(
                               CASE WHEN m.aconex_url LIKE '%MiniURL%' THEN m.aconex_url END,
                               CASE WHEN l.aconex_url LIKE '%MiniURL%' THEN l.aconex_url END,
                               'https://mea.aconex.com/authentication/index.html?Action=View&MENU_TARGET=MAIL-SEARCH&ENTITY=PROJECT&ID=1744833878&searchQuery=mailNumber%3A%22' || m.mail_no || '%22'
                           ) as aconex_url
                    FROM fact_mail m
                    LEFT JOIN mail_aconex_links l ON m.mail_no = l.mail_no
                    {where_sql}
                    ORDER BY {order_by}
                    LIMIT {limit}
                """).fetchall()
            else:
                ws.title = "RFIs Register"
                headers = ["RFI No", "Reference No", "Subject", "From Org", "To Org", "Status", "Sent Date", "Due Date", "Days Late", "Overdue", "Aconex Link"]
                filters = [
                    f"(r.project_package = '{project_id}' OR r.subject ILIKE '%{project_id}%' OR r.mail_no ILIKE '%{project_id}%')",
                    "r.from_organization NOT ILIKE '%Hill Inter%'"
                ]
                if search:
                    s_c = search.replace("'", "''")
                    filters.append(f"(r.mail_no ILIKE '%{s_c}%' OR r.subject ILIKE '%{s_c}%' OR r.from_organization ILIKE '%{s_c}%')")
                if from_org:
                    filters.append(f"r.from_organization = '{from_org.replace("'", "''")}'")
                if to_org:
                    filters.append(f"r.to_organization = '{to_org.replace("'", "''")}'")
                if mail_type:
                    filters.append(f"(r.mail_no LIKE '%-RFI-%')")
                if status:
                    filters.append(f"r.status = '{status.replace("'", "''")}'")
                if is_overdue in ['true', 'yes', '1']:
                    filters.append("r.is_overdue = true")

                order_by = "r.is_overdue DESC, r.sent_date DESC NULLS LAST"
                if sort:
                    cols_map = {'mail_no': 'r.mail_no', 'subject': 'r.subject', 'from_org': 'r.from_organization', 'to_org': 'r.to_organization', 'sent_date': 'r.sent_date', 'due_date': 'r.due_date', 'status': 'r.status'}
                    clauses = []
                    for t in sort.split(','):
                        parts = t.strip().split()
                        if parts and parts[0].lower() in cols_map:
                            d = 'DESC' if len(parts) > 1 and parts[1].upper() == 'DESC' else 'ASC'
                            clauses.append(f"{cols_map[parts[0].lower()]} {d} NULLS LAST")
                    if clauses: order_by = ", ".join(clauses)

                where_sql = f"WHERE {' AND '.join(filters)}"
                rows = con.execute(f"""
                    SELECT r.mail_no, r.mail_no, r.subject, r.from_organization, r.to_organization,
                           r.status, r.sent_date::VARCHAR, r.due_date::VARCHAR, r.days_late,
                           CASE WHEN r.is_overdue THEN 'OVERDUE' ELSE 'ON TRACK' END,
                           COALESCE(
                               CASE WHEN r.aconex_url LIKE '%MiniURL%' THEN r.aconex_url END,
                               CASE WHEN l.aconex_url LIKE '%MiniURL%' THEN l.aconex_url END,
                               'https://mea.aconex.com/authentication/index.html?Action=View&MENU_TARGET=MAIL-SEARCH&ENTITY=PROJECT&ID=1744833878&searchQuery=mailNumber%3A%22' || r.mail_no || '%22'
                           ) as aconex_url
                    FROM fact_rfis r
                    LEFT JOIN mail_aconex_links l ON r.mail_no = l.mail_no
                    {where_sql}
                    ORDER BY {order_by}
                    LIMIT {limit}
                """).fetchall()

        else: # workflows
            ws.title = "Workflows Register"
            headers = ["Workflow No", "Workflow Name", "Status", "Initiator Org", "Assigned Org", "Assigned User", "Current Step", "Overdue Steps", "Start Date", "Due Date", "Aconex Link"]
            
            filters = [f"project_package = '{project_id}'", "from_organization NOT ILIKE '%Hill Inter%'"]
            if search:
                s_c = search.replace("'", "''")
                filters.append(f"(workflow_no ILIKE '%{s_c}%' OR workflow_name ILIKE '%{s_c}%' OR initiator_organization ILIKE '%{s_c}%')")
            if status:
                filters.append(f"workflow_status = '{status.replace("'", "''")}'")
            if is_overdue in ['true', 'yes', '1']:
                filters.append("overdue_steps_count > 0")

            order_by = "overdue_steps_count DESC, start_date DESC NULLS LAST"
            if sort:
                cols_map = {'workflow_no': 'workflow_no', 'workflow_name': 'workflow_name', 'status': 'workflow_status', 'initiator': 'initiator_organization', 'current_step': 'current_step_name', 'start_date': 'start_date', 'due_date': 'original_due_date', 'overdue_steps': 'overdue_steps_count'}
                clauses = []
                for t in sort.split(','):
                    parts = t.strip().split()
                    if parts and parts[0].lower() in cols_map:
                        d = 'DESC' if len(parts) > 1 and parts[1].upper() == 'DESC' else 'ASC'
                        clauses.append(f"{cols_map[parts[0].lower()]} {d} NULLS LAST")
                if clauses: order_by = ", ".join(clauses)

            where_sql = f"WHERE {' AND '.join(filters)}"
            rows = con.execute(f"""
                SELECT workflow_no, workflow_name, workflow_status, initiator_organization, current_assigned_org,
                       current_assigned_user, current_step_name, overdue_steps_count, start_date::VARCHAR, original_due_date::VARCHAR,
                       'https://mea.aconex.com/WorkflowSearch?searchtext=' || workflow_no || '&projectid=1744833878' as aconex_url
                FROM fact_workflows
                {where_sql}
                ORDER BY {order_by}
                LIMIT {limit}
            """).fetchall()

        # 3. Build Title Banner
        num_cols = len(headers)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_cols)
        c1 = ws.cell(1, 1, f"ACONEX PMC CONTROL PLATFORM — {ws.title.upper()} — {proj_name.upper()} ({pkg_code})")
        c1.font = title_font
        c1.fill = title_fill
        c1.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[1].height = 32

        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=num_cols)
        c2 = ws.cell(2, 1, f"Exported: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | Total Exported Records: {len(rows):,} | Filters Applied: {filter_desc}")
        c2.font = sub_font
        c2.fill = sub_fill
        c2.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[2].height = 20

        # Row 3 is blank
        ws.row_dimensions[3].height = 10

        # Row 4 is Header Row
        header_row_idx = 4
        ws.row_dimensions[header_row_idx].height = 26
        for col_idx, h_text in enumerate(headers, 1):
            cell = ws.cell(header_row_idx, col_idx, h_text)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center" if "Date" in h_text or "Status" in h_text or "Link" in h_text else "left", vertical="center")
            cell.border = thin_border

        # 4. Write Data Rows
        start_data_row = 5
        for r_idx, r_data in enumerate(rows, start_data_row):
            ws.row_dimensions[r_idx].height = 20
            is_zebra = (r_idx % 2 == 0)
            row_fill = zebra_fill if is_zebra else white_fill

            for col_idx, val in enumerate(r_data, 1):
                cell = ws.cell(r_idx, col_idx)
                cell.border = thin_border
                cell.fill = row_fill

                val_str = str(val) if val is not None else ""

                # Format Aconex Hyperlink Column (last column)
                if col_idx == num_cols and val_str.startswith("http"):
                    cell.value = "Open in Aconex"
                    cell.hyperlink = val_str
                    cell.font = link_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                # Format First Column as Direct Live Hyperlink as well
                elif col_idx == 1 and len(r_data) >= num_cols and str(r_data[-1]).startswith("http"):
                    cell.value = val_str
                    cell.hyperlink = str(r_data[-1])
                    cell.font = link_font
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                # Format Overdue / Status Badges
                elif "OVERDUE" in val_str.upper():
                    cell.value = val_str
                    cell.fill = overdue_fill
                    cell.font = overdue_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif val_str in ["Closed", "Closed-Out", "Responded", "Completed"]:
                    cell.value = val_str
                    cell.fill = closed_fill
                    cell.font = closed_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif "Date" in headers[col_idx - 1]:
                    cell.value = val_str[:10] if len(val_str) >= 10 else val_str
                    cell.font = regular_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif isinstance(val, (int, float)):
                    cell.value = val
                    cell.font = bold_font if "Late" in headers[col_idx - 1] else regular_font
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    cell.value = val_str
                    cell.font = regular_font
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        # 5. Auto-fit column widths with padding
        for col_idx in range(1, num_cols + 1):
            col_letter = get_column_letter(col_idx)
            max_len = 0
            for row_idx in range(header_row_idx, min(header_row_idx + 100, len(rows) + start_data_row)):
                cell_val = str(ws.cell(row_idx, col_idx).value or "")
                if len(cell_val) > max_len:
                    max_len = len(cell_val)
            col_width = max(12, min(50, max_len + 3))
            ws.column_dimensions[col_letter].width = col_width

        # Freeze pane below header row
        ws.freeze_panes = ws['A5']

        out_buf = io.BytesIO()
        wb.save(out_buf)
        out_buf.seek(0)

        dir_tag = ""
        if cat_clean == 'letters' and direction:
            if direction.lower() == 'inbound':
                dir_tag = "_From_Contractor"
            elif direction.lower() == 'outbound':
                dir_tag = "_To_Contractor"
            elif direction.lower() == 'all':
                dir_tag = "_All_Directions"
        filename = f"{project_id}_{cat_clean.capitalize()}{dir_tag}_Log_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
        return out_buf.getvalue(), filename
    finally:
        con.close()


def generate_ball_in_court_excel_export(
    ball_in_court: str = None,
    consultant: str = None,
    file_type: str = None,
    discipline: str = None,
    package: str = None,
    process_status: str = None,
    mail_status: str = None,
    action: str = None,
    search: str = None
) -> tuple[bytes, str]:
    """
    Generates an Excel export of Ball in Court processes matching the exact
    format of 0006_PROCESS DETAILS TABLE MD2.xlsx with dual headers and styling.
    """
    con = get_db_connection()
    try:
        filters = []
        params = []
        if file_type and file_type != 'all':
            selected_types = [t.strip() for t in file_type.split(',') if t.strip()]
            if len(selected_types) == 1:
                filters.append("file_type = ?")
                params.append(selected_types[0])
            elif len(selected_types) > 1:
                placeholders = ','.join(['?' for _ in selected_types])
                filters.append(f"file_type IN ({placeholders})")
                params.extend(selected_types)

        if consultant and consultant != 'all':
            if consultant == 'DAR AL HANDASAH':
                filters.append("ball_in_court ILIKE '%DAR AL HANDASAH%'")
            elif consultant == 'OBMI':
                filters.append("ball_in_court ILIKE '%OBMI%'")
            elif consultant == 'NACO':
                filters.append("ball_in_court ILIKE '%NACO%'")
            elif consultant == 'Royal HaskoningDHV':
                filters.append("(ball_in_court ILIKE '%Haskoning%' OR ball_in_court ILIKE '%Royal%')")
            elif consultant == 'EGIS':
                filters.append("ball_in_court ILIKE '%EGIS%'")
            elif consultant == 'Aedas':
                filters.append("ball_in_court ILIKE '%Aedas%'")
            elif consultant == 'Shaker_ECG':
                filters.append("(ball_in_court ILIKE '%Shaker%' OR ball_in_court ILIKE '%ECG%')")
            elif consultant == 'Modon':
                filters.append("ball_in_court ILIKE '%Modon%'")
            elif consultant == 'Hill_Only':
                filters.append("ball_in_court = 'Hill Inter'")
            else:
                filters.append("ball_in_court ILIKE ?")
                params.append(f"%{consultant}%")

        if ball_in_court and ball_in_court != 'all':
            if ball_in_court in ['(Blanks)', 'blanks', 'blank', 'null', 'None']:
                filters.append("(ball_in_court IS NULL OR ball_in_court = '' OR ball_in_court = 'N/A')")
            elif ball_in_court == 'Hill':
                filters.append("ball_in_court ILIKE '%Hill%'")
            elif ball_in_court == 'Shaker_ECG':
                filters.append("(ball_in_court ILIKE '%Shaker%' OR ball_in_court ILIKE '%ECG%')")
            elif ball_in_court == 'Modon':
                filters.append("ball_in_court ILIKE '%Modon%'")
            elif '|||' in ball_in_court:
                selected_parties = [p.strip() for p in ball_in_court.split('|||') if p.strip()]
                placeholders = ','.join(['?' for _ in selected_parties])
                filters.append(f"ball_in_court IN ({placeholders})")
                params.extend(selected_parties)
            else:
                filters.append("ball_in_court = ?")
                params.append(ball_in_court)

        if discipline and discipline != 'all':
            filters.append("discipline = ?")
            params.append(discipline)

        if package and package != 'all':
            filters.append("project_package = ?")
            params.append(package)
        if process_status and process_status != 'all':
            filters.append("process_status = ?")
            params.append(process_status)
        if mail_status and mail_status != 'all':
            filters.append("mail_status = ?")
            params.append(mail_status)
        if action and action != 'all':
            filters.append("action = ?")
            params.append(action)
        if search:
            filters.append("(process_ref ILIKE ? OR mail_no ILIKE ? OR ball_in_court ILIKE ? OR latest_activity ILIKE ?)")
            q = f"%{search}%"
            params.extend([q, q, q, q])
        
        where_clause = f"WHERE {' AND '.join(filters)}" if filters else ""
        query = f"""
            SELECT 
                process_ref, initiated_on, process_status, duration, delay, latest_activity,
                COALESCE(discipline, 'General') as discipline,
                mail_status, ball_in_court, action, overdue_days, mail_no, aconex_url
            FROM fact_ball_in_court_processes
            {where_clause}
            ORDER BY COALESCE(overdue_days, delay, duration, 0) DESC, initiated_on DESC
        """
        rows = con.execute(query, params).fetchall()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Process Details"
        ws.views.sheetView[0].showGridLines = True

        proc_super_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        proc_super_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        
        mail_super_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        mail_super_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")

        sub_header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        sub_header_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")

        thin_side = Side(style='thin', color='CBD5E1')
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        regular_font = Font(name="Calibri", size=9, color="0F172A")
        bold_font = Font(name="Calibri", size=9, bold=True, color="0F172A")
        link_font = Font(name="Calibri", size=9, bold=True, color="0000EE", underline="single")

        overdue_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        overdue_font = Font(name="Calibri", size=9, bold=True, color="DC2626")

        outstanding_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
        outstanding_font = Font(name="Calibri", size=9, bold=True, color="D97706")

        responded_fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
        responded_font = Font(name="Calibri", size=9, bold=True, color="059669")

        # Row 1: Merged super headers
        ws.merge_cells('A1:G1')
        c1 = ws['A1']
        c1.value = "Processes"
        c1.font = proc_super_font
        c1.fill = proc_super_fill
        c1.alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells('H1:L1')
        c2 = ws['H1']
        c2.value = "Mails"
        c2.font = mail_super_font
        c2.fill = mail_super_fill
        c2.alignment = Alignment(horizontal="center", vertical="center")

        ws.row_dimensions[1].height = 24

        # Row 2: Sub headers
        sub_headers = [
            'Process Reference', 'Initiated On', 'Process Status ', 'Duration', 'Delay', 'Latest Activity',
            'Discipline', 'Mail Status', 'Ball in Court', 'Action', 'Overdue', 'Mail No.'
        ]
        for col_idx, h_text in enumerate(sub_headers, 1):
            cell = ws.cell(2, col_idx, h_text)
            cell.font = sub_header_font
            cell.fill = sub_header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border
        ws.row_dimensions[2].height = 24

        # Fill Data Rows
        for r_idx, r_data in enumerate(rows, 3):
            ws.row_dimensions[r_idx].height = 20
            p_ref, init_on, p_status, dur, delay, latest, disc, m_status, bic, action_val, ov_days, mno, url = r_data

            vals = [
                p_ref, init_on, p_status, dur, delay, latest,
                disc, m_status, bic, action_val, ov_days, mno
            ]

            for c_idx, val in enumerate(vals, 1):
                cell = ws.cell(r_idx, c_idx)
                cell.border = thin_border
                
                if c_idx == 1:
                    cell.value = val
                    if url and val:
                        cell.hyperlink = url
                        cell.font = link_font
                    else:
                        cell.font = bold_font
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                elif c_idx == 2:
                    cell.value = str(val)[:10] if val else ""
                    cell.font = regular_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif c_idx == 3:
                    cell.value = val
                    if val == 'Overdue':
                        cell.fill = overdue_fill
                        cell.font = overdue_font
                    elif val == 'Outstanding':
                        cell.fill = outstanding_fill
                        cell.font = outstanding_font
                    elif val in ('Responded', 'Closed-Out'):
                        cell.fill = responded_fill
                        cell.font = responded_font
                    else:
                        cell.font = regular_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif c_idx in (4, 5):
                    cell.value = float(val) if val is not None else ""
                    cell.font = regular_font
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    if val is not None: cell.number_format = '0.00'
                elif c_idx in (6, 7):
                    cell.value = val
                    cell.font = regular_font
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                elif c_idx == 8:
                    cell.value = val
                    if val == 'Overdue':
                        cell.fill = overdue_fill
                        cell.font = overdue_font
                    elif val == 'Outstanding':
                        cell.fill = outstanding_fill
                        cell.font = outstanding_font
                    elif val in ('Responded', 'Closed-Out'):
                        cell.fill = responded_fill
                        cell.font = responded_font
                    else:
                        cell.font = regular_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif c_idx == 9:
                    cell.value = val
                    cell.font = bold_font
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                elif c_idx == 10:
                    cell.value = val
                    if val == 'To Respond':
                        cell.font = overdue_font
                    else:
                        cell.font = regular_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif c_idx == 11:
                    cell.value = float(val) if val is not None else ""
                    cell.font = overdue_font if val else regular_font
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    if val is not None: cell.number_format = '0.00'
                elif c_idx == 12:
                    cell.value = val
                    if url and val:
                        cell.hyperlink = url
                        cell.font = link_font
                    else:
                        cell.font = bold_font
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        for col_idx in range(1, 13):
            col_letter = get_column_letter(col_idx)
            max_len = 0
            for row_idx in range(2, min(50, len(rows) + 3)):
                v = str(ws.cell(row_idx, col_idx).value or "")
                if len(v) > max_len: max_len = len(v)
            ws.column_dimensions[col_letter].width = max(12, min(45, max_len + 3))

        ws.freeze_panes = ws['A3']

        out_buf = io.BytesIO()
        wb.save(out_buf)
        out_buf.seek(0)
        filename = f"Ball_in_Court_Processes_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
        return out_buf.getvalue(), filename
    finally:
        con.close()


def generate_transmittals_excel_export(
    project_id: str,
    search: str = None,
    doc_type: str = None,
    discipline: str = None,
    status: str = None
) -> tuple[bytes, str]:
    """
    Generates an Excel export of Document Transmittals for a project package,
    matching the user's layout and styling.
    """
    con = get_db_connection()
    try:
        filters = ["project_package = ?"]
        params = [project_id]

        if search and isinstance(search, str):
            clean_s = search.replace("'", "''")
            filters.append(f"(document_no ILIKE '%{clean_s}%' OR title ILIKE '%{clean_s}%' OR transmittal_no ILIKE '%{clean_s}%' OR transmitted_to_org ILIKE '%{clean_s}%' OR transmittal_in_from_org ILIKE '%{clean_s}%' OR transmitted_recipients ILIKE '%{clean_s}%')")

        if doc_type and doc_type != 'all':
            if ',' in doc_type:
                types = [t.strip().replace("'", "''") for t in doc_type.split(',') if t.strip()]
                t_str = "','".join(types)
                filters.append(f"doc_type IN ('{t_str}')")
            else:
                filters.append(f"doc_type = '{doc_type.replace("'", "''")}'")

        if discipline and discipline != 'all':
            filters.append(f"discipline = '{discipline.replace("'", "''")}'")

        if status and status != 'all':
            filters.append(f"doc_status = '{status.replace("'", "''")}'")

        where_clause = f"WHERE {' AND '.join(filters)}"
        query = f"""
            SELECT 
                document_no, title, doc_type, discipline, doc_status, revision,
                date_modified, transmittal_no, transmitted_sent_date,
                transmitted_to_org, transmitted_recipients, reason_for_issue,
                transmittal_in_from_org, aconex_doc_url, aconex_transmittal_url
            FROM fact_document_transmittals
            {where_clause}
            ORDER BY date_modified DESC NULLS LAST
            LIMIT 50000
        """
        rows = con.execute(query, params).fetchall()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Transmittals History"
        ws.views.sheetView[0].showGridLines = True

        header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        thin_side = Side(style='thin', color='CBD5E1')
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        regular_font = Font(name="Calibri", size=9, color="0F172A")
        bold_font = Font(name="Calibri", size=9, bold=True, color="0F172A")
        link_font = Font(name="Calibri", size=9, bold=True, color="059669", underline="single")

        headers = [
            'Document No', 'Document Title', 'Document Type', 'Discipline', 'Status',
            'Rev', 'Modified Date', 'Transmittal No', 'Sent Date', 'To Org',
            'Recipients', 'Reason For Issue', 'From Org'
        ]

        for col_idx, h_text in enumerate(headers, 1):
            cell = ws.cell(1, col_idx, h_text)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border
        ws.row_dimensions[1].height = 26

        for r_idx, r in enumerate(rows, 2):
            ws.row_dimensions[r_idx].height = 20
            doc_no, title, dtype, disc, st, rev, dt_mod, trans_no, sent_dt, to_org, recip, reason, from_org, doc_url, trans_url = r

            vals = [
                doc_no, title, dtype, disc, st,
                rev, dt_mod, trans_no, sent_dt, to_org,
                recip, reason, from_org
            ]

            for c_idx, val in enumerate(vals, 1):
                cell = ws.cell(r_idx, c_idx)
                cell.border = thin_border
                cell.value = val

                if c_idx == 1:
                    if doc_url and val:
                        cell.hyperlink = doc_url
                        cell.font = link_font
                    else:
                        cell.font = bold_font
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                elif c_idx == 8:
                    if trans_url and val and val != 'N/A':
                        cell.hyperlink = trans_url
                        cell.font = link_font
                    else:
                        cell.font = bold_font
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                elif c_idx in (5, 6):
                    cell.font = bold_font if c_idx == 5 else regular_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.font = regular_font
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        for col_idx in range(1, len(headers) + 1):
            col_letter = get_column_letter(col_idx)
            max_len = 0
            for row_idx in range(1, min(50, len(rows) + 2)):
                v = str(ws.cell(row_idx, col_idx).value or "")
                if len(v) > max_len: max_len = len(v)
            ws.column_dimensions[col_letter].width = max(12, min(50, max_len + 3))

        ws.freeze_panes = ws['A2']

        out_buf = io.BytesIO()
        wb.save(out_buf)
        out_buf.seek(0)
        filename = f"Transmittals_{project_id}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
        return out_buf.getvalue(), filename
    finally:
        con.close()


