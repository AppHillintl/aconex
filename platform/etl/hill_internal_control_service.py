import io
import os
import re
import urllib.parse
from datetime import datetime
from email.message import EmailMessage
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import duckdb

import sys
backend_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "backend"))
if backend_path not in sys.path:
    sys.path.insert(0, backend_path)

import package_resolver as pr

from pathlib import Path as _Path
_SRV_ROOT = _Path(__file__).resolve().parent.parent.parent
_SRV_CLOUD_DB = _SRV_ROOT / "aconex_cloud.duckdb"
_SRV_LOCAL_DB = _SRV_ROOT / "aconex.duckdb"

if _SRV_LOCAL_DB.exists():
    DB_PATH = str(_SRV_LOCAL_DB).replace("\\", "/")
elif _SRV_CLOUD_DB.exists():
    DB_PATH = str(_SRV_CLOUD_DB).replace("\\", "/")
else:
    DB_PATH = str(_SRV_LOCAL_DB).replace("\\", "/")

def get_db():
    return duckdb.connect(DB_PATH, read_only=True)

def get_hill_directory_users() -> dict:
    """Loads all Hill International engineers and their verified contact info from dim_user_directory."""
    con = get_db()
    try:
        rows = con.execute("""
            SELECT user_name, clean_name, email, position
            FROM dim_user_directory
            WHERE organization_name ILIKE '%Hill%' AND email IS NOT NULL AND email != ''
        """).fetchall()
        user_map = {}
        for r in rows:
            uname, cname, email, pos = r
            user_map[cname.lower().strip()] = {
                "clean_name": cname,
                "user_name": uname,
                "email": email,
                "position": pos or "PMC Engineer"
            }
        return user_map
    finally:
        con.close()

def match_hill_engineer(raw_name: str, user_map: dict) -> dict:
    """Matches an assignee string against Hill directory with title stripping."""
    clean_p = re.sub(r'^(Mr|Ms|Eng)\s+', '', raw_name, flags=re.I).strip()
    if not clean_p or 'DC' in clean_p.upper():
        return None
    p_lower = clean_p.lower()
    if p_lower in user_map:
        return user_map[p_lower]
    for k, info in user_map.items():
        if p_lower in k or k in p_lower:
            return info
    return {
        "clean_name": clean_p,
        "user_name": raw_name,
        "email": pr.resolve_user_email(clean_p, "Hill International"),
        "position": "PMC Engineer"
    }

def get_hill_team_summary(package_filter: str = None) -> list:
    """
    Returns aggregated direct-action backlog for all Hill International engineers.
    Strictly includes only direct assignees for active/overdue review steps and direct overdue mails.
    Strictly excludes CC and FYI items.
    """
    con = get_db()
    user_map = get_hill_directory_users()
    try:
        where_wf = "WHERE s.assigned_to_organization ILIKE '%Hill%' AND s.step_status IN ('Overdue', 'Current')"

        steps = con.execute(f"""
            SELECT 
                s.workflow_no,
                s.step_name,
                s.step_status,
                s.assigned_to_name,
                s.days_late,
                s.is_overdue,
                COALESCE(w.workflow_name, s.workflow_no),
                s.document_no,
                s.document_title
            FROM fact_workflow_steps s
            LEFT JOIN fact_workflows w ON s.workflow_no = w.workflow_no
            {where_wf}
        """).fetchall()

        # Mails direct to Hill users
        where_mail = "WHERE to_organization ILIKE '%Hill%' AND status NOT IN ('Closed', 'Responded') AND is_overdue = true"
        if package_filter and package_filter != 'all':
            clean_pkg = package_filter.replace("'", "''")
            where_mail += f" AND project_package = '{clean_pkg}'"

        mails = con.execute(f"""
            SELECT mail_no, mail_type, subject, recipients, days_late, project_package
            FROM fact_mail
            {where_mail}
        """).fetchall()

        engineers = {}

        # 1. Process review steps
        for row in steps:
            wf_no, s_name, s_status, raw_assignees, days_late, is_overdue, wf_name, doc_no, doc_title = row
            pkg = pr.resolve_project_package(doc_no=doc_no, title=doc_title)
            if package_filter and package_filter != 'all' and pkg != package_filter:
                continue
            parts = [p.strip() for p in (raw_assignees or "").split(',') if p.strip()]
            for p in parts:
                matched = match_hill_engineer(p, user_map)
                if not matched:
                    continue
                cname = matched["clean_name"]
                if cname not in engineers:
                    engineers[cname] = {
                        "name": cname,
                        "email": matched["email"],
                        "position": matched["position"],
                        "overdue_steps": 0,
                        "current_steps": 0,
                        "total_steps": 0,
                        "overdue_mails": 0,
                        "total_items": 0,
                        "max_delay_days": 0,
                        "workflows_set": set(),
                        "packages_set": set()
                    }
                eng = engineers[cname]
                eng["workflows_set"].add(wf_no)
                if pkg:
                    eng["packages_set"].add(pkg)
                d_late = days_late or 0
                if d_late > eng["max_delay_days"]:
                    eng["max_delay_days"] = d_late
                
                eng["total_steps"] += 1
                if is_overdue or s_status == 'Overdue':
                    eng["overdue_steps"] += 1
                else:
                    eng["current_steps"] += 1

        # 2. Process direct mails
        for m in mails:
            m_no, m_type, m_subj, raw_recipients, days_late, pkg = m
            parts = [p.strip() for p in (raw_recipients or "").split(';') if p.strip()]
            for p in parts:
                matched = match_hill_engineer(p, user_map)
                if not matched:
                    continue
                cname = matched["clean_name"]
                if cname not in engineers:
                    engineers[cname] = {
                        "name": cname,
                        "email": matched["email"],
                        "position": matched["position"],
                        "overdue_steps": 0,
                        "current_steps": 0,
                        "total_steps": 0,
                        "overdue_mails": 0,
                        "total_items": 0,
                        "max_delay_days": 0,
                        "workflows_set": set(),
                        "packages_set": set()
                    }
                eng = engineers[cname]
                if pkg:
                    eng["packages_set"].add(pkg)
                d_late = days_late or 0
                if d_late > eng["max_delay_days"]:
                    eng["max_delay_days"] = d_late
                eng["overdue_mails"] += 1

        # Finalize counts
        result = []
        for eng in engineers.values():
            eng["total_workflows"] = len(eng["workflows_set"])
            eng["packages"] = sorted(list(eng["packages_set"]))
            eng["total_items"] = eng["total_steps"] + eng["overdue_mails"]
            eng["total_overdue"] = eng["overdue_steps"] + eng["overdue_mails"]
            del eng["workflows_set"]
            del eng["packages_set"]
            result.append(eng)

        result.sort(key=lambda x: (x["total_overdue"], x["total_items"]), reverse=True)
        return result
    finally:
        con.close()

def get_hill_engineer_dossier(engineer_name: str, package_filter: str = None) -> dict:
    """
    Returns the comprehensive, direct-action portfolio for a single Hill engineer.
    Includes itemized review steps, direct mails, SLA metrics, and ready-to-copy HTML reminder email.
    """
    con = get_db()
    user_map = get_hill_directory_users()
    matched_self = match_hill_engineer(engineer_name, user_map)
    clean_target = matched_self["clean_name"] if matched_self else engineer_name.strip()
    target_lower = clean_target.lower()

    try:
        where_wf = "WHERE s.assigned_to_organization ILIKE '%Hill%' AND s.step_status IN ('Overdue', 'Current')"

        steps = con.execute(f"""
            SELECT 
                s.workflow_no,
                s.step_name,
                s.step_status,
                s.assigned_to_name,
                s.document_no,
                s.document_title,
                s.days_late,
                s.is_overdue,
                COALESCE(w.workflow_name, s.workflow_no),
                w.start_date::VARCHAR,
                w.original_due_date::VARCHAR,
                s.date_in::VARCHAR,
                s.date_due::VARCHAR,
                'https://mea.aconex.com/WorkflowSearch?searchtext=' || s.workflow_no || '&projectid=1744833878' AS aconex_url,
                CASE 
                    WHEN d.document_id IS NOT NULL AND LENGTH(TRIM(d.document_id)) > 3 AND d.document_id != 'None'
                    THEN 'https://mea.aconex.com/ViewDoc?trackingid=' || d.document_id || '&projectid=1744833878&docno=' || s.document_no || '&username=aabdelrahim4&action=ViewDocumentSummary'
                    ELSE 'https://mea.aconex.com/authentication/index.html?Action=View&MENU_TARGET=DOC-SEARCH&ENTITY=PROJECT&ID=1744833878&searchQuery=documentNumber%3A%22' || COALESCE(s.document_no, '') || '%22'
                END AS doc_aconex_url
            FROM fact_workflow_steps s
            LEFT JOIN fact_workflows w ON s.workflow_no = w.workflow_no
            LEFT JOIN (SELECT document_no, MAX(tracking_id) as document_id, MAX(aconex_url) as doc_direct_url FROM dim_documents GROUP BY document_no) d ON s.document_no = d.document_no
            {where_wf}
            ORDER BY s.is_overdue DESC, s.days_late DESC NULLS LAST, s.date_in DESC
        """).fetchall()

        matched_steps = []
        for r in steps:
            doc_no = r[4]
            doc_title = r[5]
            pkg = pr.resolve_project_package(doc_no=doc_no, title=doc_title)
            if package_filter and package_filter != 'all' and pkg != package_filter:
                continue
            raw_assignees = r[3] or ""
            parts = [p.strip() for p in raw_assignees.split(',') if p.strip()]
            is_hit = False
            for p in parts:
                m = match_hill_engineer(p, user_map)
                if m and m["clean_name"].lower() == target_lower:
                    is_hit = True
                    break
            if is_hit:
                matched_steps.append({
                    "workflow_no": r[0],
                    "step_name": r[1],
                    "step_status": r[2],
                    "assigned_to": r[3],
                    "document_no": r[4],
                    "document_title": r[5],
                    "days_late": r[6] or 0,
                    "is_overdue": r[7],
                    "workflow_name": r[8],
                    "package": pkg or "Unassigned",
                    "start_date": r[9][:10] if r[9] else None,
                    "due_date": r[10][:10] if r[10] else (r[12][:10] if r[12] else None),
                    "date_due": r[10][:10] if r[10] else (r[12][:10] if r[12] else None),
                    "date_in": r[11][:10] if r[11] else None,
                    "aconex_url": r[13],
                    "doc_aconex_url": r[14] if len(r) > 14 else r[13]
                })

        # Direct Mails
        where_mail = "WHERE to_organization ILIKE '%Hill%' AND status NOT IN ('Closed', 'Responded') AND is_overdue = true"
        if package_filter and package_filter != 'all':
            clean_pkg = package_filter.replace("'", "''")
            where_mail += f" AND project_package = '{clean_pkg}'"

        mails = con.execute(f"""
            SELECT 
                mail_no, mail_type, subject, reference_number, from_organization,
                recipients, status, sent_date::VARCHAR, due_date::VARCHAR, days_late, project_package,
                COALESCE(aconex_url, 'https://mea.aconex.com/authentication/index.html?Action=View&MENU_TARGET=MAIL-SEARCH&ENTITY=PROJECT&ID=1744833878&searchQuery=mailNumber%3A%22' || mail_no || '%22') AS aconex_url
            FROM fact_mail
            {where_mail}
            ORDER BY days_late DESC NULLS LAST, sent_date DESC
        """).fetchall()

        matched_mails = []
        for m in mails:
            raw_recipients = m[5] or ""
            parts = [p.strip() for p in raw_recipients.split(';') if p.strip()]
            is_hit = False
            for p in parts:
                m_info = match_hill_engineer(p, user_map)
                if m_info and m_info["clean_name"].lower() == target_lower:
                    is_hit = True
                    break
            if is_hit:
                matched_mails.append({
                    "mail_no": m[0],
                    "mail_type": m[1],
                    "subject": m[2],
                    "ref_no": m[3],
                    "from_org": m[4],
                    "recipients": m[5],
                    "status": m[6],
                    "sent_date": m[7][:10] if m[7] else None,
                    "due_date": m[8][:10] if m[8] else None,
                    "days_late": m[9] or 0,
                    "package": m[10] or "Unassigned",
                    "aconex_url": m[11]
                })

        total_overdue = sum(1 for s in matched_steps if s["is_overdue"]) + len(matched_mails)
        total_current = sum(1 for s in matched_steps if not s["is_overdue"])
        pkgs = sorted(list(set([s["package"] for s in matched_steps] + [m["package"] for m in matched_mails])))

        # Generate HTML email body
        email_html = generate_hill_reminder_html(
            engineer_name=clean_target,
            email=matched_self["email"] if matched_self else "",
            position=matched_self["position"] if matched_self else "PMC Engineer",
            steps=matched_steps,
            mails=matched_mails,
            pkgs=pkgs
        )

        return {
            "engineer": {
                "name": clean_target,
                "email": matched_self["email"] if matched_self else "",
                "position": matched_self["position"] if matched_self else "PMC Engineer",
            },
            "summary": {
                "total_items": len(matched_steps) + len(matched_mails),
                "total_overdue": total_overdue,
                "total_current": total_current,
                "total_steps": len(matched_steps),
                "total_mails": len(matched_mails),
                "packages": pkgs
            },
            "steps": matched_steps,
            "mails": matched_mails,
            "email_html": email_html,
            "subject": f"URGENT: Hill PMC Direct Action Items Backlog - {clean_target} ({total_overdue} Overdue)"
        }
    finally:
        con.close()

def generate_hill_reminder_html(engineer_name: str, email: str, position: str, steps: list, mails: list, pkgs: list) -> str:
    """Generates an executive, branded HTML email for Hill International internal reminders."""
    overdue_steps = [s for s in steps if s["is_overdue"]]
    current_steps = [s for s in steps if not s["is_overdue"]]
    total_ov = len(overdue_steps) + len(mails)
    pkg_str = ", ".join(pkgs) if pkgs else "All Packages"

    html = f"""<!DOCTYPE html>
<html>
<head>
<meta charset="UTF-8">
<style>
body {{ font-family: 'Segoe UI', Arial, sans-serif; background-color: #f8fafc; color: #1e293b; margin: 0; padding: 24px; }}
.container {{ max-width: 900px; margin: 0 auto; background: #ffffff; border: 1px solid #e2e8f0; border-radius: 12px; overflow: hidden; box-shadow: 0 4px 6px -1px rgba(0,0,0,0.05); }}
.header {{ background: linear-gradient(135deg, #0f172a 0%, #1e293b 100%); color: #ffffff; padding: 24px 32px; border-bottom: 4px solid #e11d48; }}
.title {{ font-size: 20px; font-weight: 800; margin: 0 0 6px 0; letter-spacing: -0.5px; }}
.subtitle {{ font-size: 13px; color: #94a3b8; margin: 0; }}
.kpi-row {{ display: flex; background: #f1f5f9; padding: 16px 32px; border-bottom: 1px solid #e2e8f0; gap: 20px; }}
.kpi-card {{ flex: 1; background: #ffffff; padding: 12px 16px; border-radius: 8px; border: 1px solid #cbd5e1; }}
.kpi-val {{ font-size: 22px; font-weight: 800; font-family: Consolas, monospace; margin-top: 4px; }}
.content {{ padding: 28px 32px; }}
h3 {{ font-size: 15px; font-weight: 700; color: #0f172a; margin: 24px 0 12px 0; border-left: 4px solid #0284c7; padding-left: 10px; }}
table {{ width: 100%; border-collapse: collapse; font-size: 12px; margin-bottom: 24px; }}
th {{ background: #0f172a; color: #ffffff; text-align: left; padding: 10px 12px; font-weight: 600; text-transform: uppercase; font-size: 11px; }}
td {{ padding: 10px 12px; border-bottom: 1px solid #e2e8f0; }}
tr:hover {{ background-color: #f8fafc; }}
.badge-late {{ background: #ffe4e6; color: #be123c; font-weight: 700; padding: 2px 8px; border-radius: 9999px; font-size: 10px; font-family: monospace; display: inline-block; }}
.badge-cur {{ background: #e0f2fe; color: #0369a1; font-weight: 700; padding: 2px 8px; border-radius: 9999px; font-size: 10px; font-family: monospace; display: inline-block; }}
.aconex-link {{ color: #0284c7; text-decoration: none; font-weight: 700; font-family: monospace; }}
.aconex-link:hover {{ text-decoration: underline; }}
.footer {{ background: #f8fafc; padding: 20px 32px; border-top: 1px solid #e2e8f0; font-size: 11px; color: #64748b; line-height: 1.6; }}
</style>
</head>
<body>
<div class="container">
    <div class="header">
        <div class="title">Hill International PMC | Direct Action Backlog Notification</div>
        <div class="subtitle">Ras El Hekma Project Controls &mdash; Attention: <strong>{engineer_name}</strong> ({position})</div>
    </div>
    <div class="kpi-row">
        <div class="kpi-card">
            <div style="font-size: 11px; color: #64748b; font-weight: 600;">TOTAL OVERDUE</div>
            <div class="kpi-val" style="color: #e11d48;">{total_ov} Items</div>
        </div>
        <div class="kpi-card">
            <div style="font-size: 11px; color: #64748b; font-weight: 600;">ACTIVE IN REVIEW</div>
            <div class="kpi-val" style="color: #0284c7;">{len(current_steps)} Steps</div>
        </div>
        <div class="kpi-card">
            <div style="font-size: 11px; color: #64748b; font-weight: 600;">ASSIGNED PACKAGES</div>
            <div class="kpi-val" style="color: #334155; font-size: 14px; margin-top: 8px;">{pkg_str}</div>
        </div>
    </div>
    <div class="content">
        <p style="font-size: 13px; line-height: 1.6; color: #334155; margin-top: 0;">
            Dear {engineer_name},<br><br>
            The following table summarizes the pending deliverables currently assigned directly to your action in Oracle Aconex. Please note that this reminder strictly contains <strong>direct review steps and technical correspondences</strong> requiring your review or sign-off (CC/FYI transmittals are strictly excluded).
        </p>
"""

    if overdue_steps:
        html += f"""
        <h3>⚠️ Overdue Review Workflow Steps ({len(overdue_steps)})</h3>
        <table>
            <thead>
                <tr>
                    <th style="width: 13%;">Workflow No</th>
                    <th style="width: 16%;">Document No</th>
                    <th style="width: 10%;">Package</th>
                    <th style="width: 16%;">Step Name</th>
                    <th style="width: 25%;">Workflow / Submittal Title</th>
                    <th style="width: 10%;">Due Date</th>
                    <th style="width: 10%;">Delay</th>
                </tr>
            </thead>
            <tbody>
        """
        for s in overdue_steps:
            html += f"""
                <tr>
                    <td><a href="{s['aconex_url']}" target="_blank" class="aconex-link">{s['workflow_no']}</a></td>
                    <td><a href="{s.get('doc_aconex_url', s['aconex_url'])}" target="_blank" class="aconex-link">{s.get('document_no') or '-'}</a></td>
                    <td><strong style="color: #0369a1;">{s['package']}</strong></td>
                    <td style="font-weight: 600;">{s['step_name']}</td>
                    <td>{s.get('document_title') or s['workflow_name']}</td>
                    <td style="font-family: monospace;">{s['due_date'] or '-'}</td>
                    <td><span class="badge-late">+{s['days_late']} days</span></td>
                </tr>
            """
        html += "</tbody></table>"

    if mails:
        html += f"""
        <h3>📬 Overdue Direct Correspondence & RFIs ({len(mails)})</h3>
        <table>
            <thead>
                <tr>
                    <th style="width: 18%;">Mail No</th>
                    <th style="width: 10%;">Package</th>
                    <th style="width: 14%;">Type</th>
                    <th style="width: 34%;">Subject</th>
                    <th style="width: 12%;">From Org</th>
                    <th style="width: 12%;">Delay</th>
                </tr>
            </thead>
            <tbody>
        """
        for m in mails:
            html += f"""
                <tr>
                    <td><a href="{m['aconex_url']}" target="_blank" class="aconex-link">{m['mail_no']}</a></td>
                    <td><strong style="color: #0369a1;">{m['package']}</strong></td>
                    <td>{m['mail_type']}</td>
                    <td>{m['subject']}</td>
                    <td>{m['from_org']}</td>
                    <td><span class="badge-late">+{m['days_late']} days</span></td>
                </tr>
            """
        html += "</tbody></table>"

    if current_steps:
        html += f"""
        <h3>⏳ Active Review Workflow Steps Within SLA ({len(current_steps)})</h3>
        <table>
            <thead>
                <tr>
                    <th style="width: 14%;">Workflow No</th>
                    <th style="width: 18%;">Document No</th>
                    <th style="width: 10%;">Package</th>
                    <th style="width: 18%;">Step Name</th>
                    <th style="width: 28%;">Workflow Title</th>
                    <th style="width: 12%;">Due Date</th>
                </tr>
            </thead>
            <tbody>
        """
        for s in current_steps:
            html += f"""
                <tr>
                    <td><a href="{s['aconex_url']}" target="_blank" class="aconex-link">{s['workflow_no']}</a></td>
                    <td><a href="{s.get('doc_aconex_url', s['aconex_url'])}" target="_blank" class="aconex-link">{s.get('document_no') or '-'}</a></td>
                    <td><strong style="color: #0369a1;">{s['package']}</strong></td>
                    <td style="font-weight: 600;">{s['step_name']}</td>
                    <td>{s.get('document_title') or s['workflow_name']}</td>
                    <td style="font-family: monospace;">{s['due_date'] or '-'}</td>
                </tr>
            """
        html += "</tbody></table>"

    html += f"""
        <p style="font-size: 13px; color: #475569; margin-bottom: 0;">
            Please access Oracle Aconex using the direct hyperlinks above and complete your reviews to ensure contractual turnaround times are maintained.
        </p>
    </div>
    <div class="footer">
        Ras El Hekma Urban Development Project Controls &mdash; Hill International PMC<br>
        Automated Governance Engine &bull; Generated on {datetime.now().strftime('%d %b %Y, %H:%M')}
    </div>
</div>
</body>
</html>
"""
    return html

def generate_hill_engineer_excel(engineer_name: str, package_filter: str = None) -> tuple[bytes, str]:
    """Generates an executive, professionally styled Excel register for the Hill engineer's action backlog."""
    dossier = get_hill_engineer_dossier(engineer_name, package_filter)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Direct Action Backlog"

    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    late_fill = PatternFill(start_color="FFE4E6", end_color="FFE4E6", fill_type="solid")
    late_font = Font(name="Arial", size=10, bold=True, color="BE123C")

    # Title block
    ws.append(["Hill International PMC - Personal Action Dossier"])
    ws.cell(1, 1).font = Font(name="Arial", size=14, bold=True, color="0F172A")
    ws.append([f"Engineer: {dossier['engineer']['name']} ({dossier['engineer']['position']})", f"Email: {dossier['engineer']['email']}"])
    ws.append([f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}", f"Total Items: {dossier['summary']['total_items']} ({dossier['summary']['total_overdue']} Overdue)"])
    ws.append([])

    link_font = Font(name="Arial", size=10, bold=True, color="0284C7", underline="single")

    # Table Header
    headers = ["Deliverable No", "Document / Ref No", "Item Type", "Package", "Step / Status", "Subject / Title", "Date In", "Due Date", "Days Late", "Aconex Link"]
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        c = ws.cell(5, col)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")

    row_idx = 6
    # Workflow Steps
    for s in dossier["steps"]:
        ws.append([
            s["workflow_no"],
            s.get("document_no") or "-",
            "Workflow Review",
            s["package"],
            s["step_name"],
            s.get("document_title") or s["workflow_name"],
            s["date_in"] or "",
            s["due_date"] or "",
            s["days_late"] if s["is_overdue"] else 0,
            "Open in Aconex"
        ])
        c1 = ws.cell(row_idx, 1)
        c1.hyperlink = s["aconex_url"]
        c1.font = link_font

        c2 = ws.cell(row_idx, 2)
        if s.get("document_no") and s.get("document_no") != "-":
            c2.hyperlink = s.get("doc_aconex_url", s["aconex_url"])
            c2.font = link_font

        c10 = ws.cell(row_idx, 10)
        c10.hyperlink = s.get("doc_aconex_url", s["aconex_url"])
        c10.font = link_font

        if s["is_overdue"]:
            ws.cell(row_idx, 9).fill = late_fill
            ws.cell(row_idx, 9).font = late_font
        row_idx += 1

    # Mails
    for m in dossier["mails"]:
        ws.append([
            m["mail_no"],
            m.get("ref_no") or m["mail_no"],
            m["mail_type"],
            m["package"],
            m["status"],
            m["subject"],
            m["sent_date"] or "",
            m["due_date"] or "",
            m["days_late"],
            "Open in Aconex"
        ])
        c1 = ws.cell(row_idx, 1)
        c1.hyperlink = m["aconex_url"]
        c1.font = link_font

        c10 = ws.cell(row_idx, 10)
        c10.hyperlink = m["aconex_url"]
        c10.font = link_font

        ws.cell(row_idx, 9).fill = late_fill
        ws.cell(row_idx, 9).font = late_font
        row_idx += 1

    # Column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    filename = f"Hill_Action_Dossier_{re.sub(r'[^a-zA-Z0-9_]', '_', dossier['engineer']['name'])}.xlsx"
    return out.getvalue(), filename

def generate_hill_engineer_eml(engineer_name: str, package_filter: str = None) -> tuple[bytes, str]:
    """Generates an Outlook-compatible .eml draft pre-attached with the Excel register."""
    dossier = get_hill_engineer_dossier(engineer_name, package_filter)
    excel_bytes, excel_name = generate_hill_engineer_excel(engineer_name, package_filter)

    msg = EmailMessage()
    msg["Subject"] = dossier["subject"]
    msg["From"] = "Ras El Hekma PMC Controls <pmc-controls@hillintl.com>"
    target_email = dossier["engineer"]["email"] or "pmc-controls@hillintl.com"
    msg["To"] = target_email
    msg["CC"] = "DeanEnglish@hillintl.com"
    msg["X-Unsent"] = "1"

    msg.set_content(
        f"Dear {dossier['engineer']['name']},\n\n"
        f"Please find attached your direct action portfolio ({dossier['summary']['total_overdue']} Overdue items).\n"
        f"Please review the attached register and complete pending actions in Aconex.\n"
    )
    msg.add_alternative(dossier["email_html"], subtype="html")

    msg.add_attachment(
        excel_bytes,
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=excel_name
    )

    clean_name = re.sub(r'[^a-zA-Z0-9_]', '_', dossier['engineer']['name'])
    eml_filename = f"Hill_Action_Reminder_{clean_name}.eml"
    return msg.as_bytes(), eml_filename
