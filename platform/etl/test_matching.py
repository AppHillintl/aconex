import openpyxl, datetime
from python_calamine import CalamineWorkbook

wb_log = openpyxl.load_workbook(r'C:\Users\bimst\Downloads\ACONEX\Data\update\Up\Logics.xlsx')
ws = wb_log['projecrt']

rules = []
for r in range(2, ws.max_row + 1):
    rules.append({
        'pkg_no': ws.cell(r, 1).value,
        'contract_no': ws.cell(r, 2).value,
        'pkg_name': ws.cell(r, 3).value,
        'contractor': ws.cell(r, 4).value,
        'project': ws.cell(r, 5).value,
        'rfi_rule': ws.cell(r, 6).value or '',
        'letter_rule': ws.cell(r, 7).value or ''
    })

def parse_clause(clause_str):
    clause_str = clause_str.strip()
    if not clause_str:
        return None
    for prefix in ['Subject ', 'From Organization ', 'From ', 'Recipients ', 'Type ']:
        if clause_str.startswith(prefix):
            return (prefix.strip(), clause_str[len(prefix):].strip())
    return ('Unknown', clause_str)

def compile_rules(rule_str):
    if not rule_str:
        return []
    return [parse_clause(p) for p in rule_str.split(chr(215)) if parse_clause(p)]

def match_clause(clause, subject, from_org, from_name, recipients, mail_type):
    field, val = clause
    val_lower = val.lower()
    if field == 'Subject':
        subj_l = (subject or '').lower()
        if ' or ' in val_lower:
            return any(opt.strip() in subj_l for opt in val_lower.split(' or ') if opt.strip())
        return val_lower in subj_l
    elif field == 'From Organization':
        org_l = (from_org or '').lower().replace('\u2013', '-').replace('–', '-')
        val_clean = val_lower.replace('\u2013', '-').replace('–', '-')
        if 'trojan' in val_clean:
            return 'trojan' in org_l or 'cox' in org_l or ('orascom' in org_l and 'trojan' in org_l)
        return val_clean in org_l
    elif field == 'From':
        return val_lower in (from_name or '').lower()
    elif field == 'Recipients':
        return val_lower in (recipients or '').lower()
    elif field == 'Type':
        type_l = (mail_type or '').lower()
        if ',' in val_lower:
            return any(opt.strip() in type_l for opt in val_lower.split(',') if opt.strip())
        return val_lower in type_l
    return False

compiled = []
for r in rules:
    compiled.append({
        'project': r['project'],
        'rfi_clauses': compile_rules(r['rfi_rule']),
        'letter_clauses': compile_rules(r['letter_rule'])
    })

wb_data = CalamineWorkbook.from_path(r'C:\Users\bimst\Downloads\ACONEX\Data\update\Up\RFI-Letters.xlsx')
rows = wb_data.get_sheet_by_name('Mail').to_python()
headers = rows[9]
mail_no_idx = headers.index('Mail No')
subj_idx = headers.index('Subject')
from_idx = headers.index('From')
org_idx = headers.index('From Organization')
recip_idx = headers.index('Recipients')
type_idx = headers.index('Type')
due_idx = headers.index('Due')
closed_idx = headers.index('Closed Out Date')
parent_st_idx = headers.index('Parent Status')

data_rows = rows[10:]
today = datetime.date(2026, 10, 12)

stats = {r['project']: {'rfi_total': 0, 'rfi_overdue': 0, 'let_total': 0, 'let_overdue': 0} for r in compiled}

for row in data_rows:
    subj = row[subj_idx] or ''
    m_from = row[from_idx] or ''
    m_org = row[org_idx] or ''
    m_recip = row[recip_idx] or ''
    m_type = row[type_idx] or ''
    due_str = str(row[due_idx]).strip() if row[due_idx] else ''
    closed_str = str(row[closed_idx]).strip() if row[closed_idx] else ''
    parent_st = str(row[parent_st_idx]).strip() if row[parent_st_idx] else ''
    
    is_ov = False
    if due_str:
        try:
            parts = due_str.split('/')
            d = datetime.date(int(parts[2]), int(parts[1]), int(parts[0]))
            if d < today and not closed_str and parent_st not in ('Responded', 'Closed-Out'):
                is_ov = True
        except Exception:
            pass

    for r in compiled:
        if m_type == 'Request For Information':
            if all(match_clause(c, subj, m_org, m_from, m_recip, m_type) for c in r['rfi_clauses']):
                stats[r['project']]['rfi_total'] += 1
                if is_ov: stats[r['project']]['rfi_overdue'] += 1
                break
        else:
            if all(match_clause(c, subj, m_org, m_from, m_recip, m_type) for c in r['letter_clauses']):
                stats[r['project']]['let_total'] += 1
                if is_ov: stats[r['project']]['let_overdue'] += 1
                break

print('=== PROJECT STATS ACCORDING TO LOGICS.XLSX ===')
for p, s in stats.items():
    rf_tot = s['rfi_total']
    rf_ov = s['rfi_overdue']
    lt_tot = s['let_total']
    lt_ov = s['let_overdue']
    tot = rf_tot + lt_tot
    tot_ov = rf_ov + lt_ov
    print(f"{p:<12} | RFI: {rf_tot:>3} (ov: {rf_ov:>3}) | Letters: {lt_tot:>3} (ov: {lt_ov:>3}) | Total: {tot:>3} (ov: {tot_ov:>3})")
