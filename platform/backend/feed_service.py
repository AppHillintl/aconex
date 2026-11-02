"""
=============================================================================
DAILY ACTIVITY FEED SERVICE (Unified Today & Yesterday Deliverables Engine)
Aggregates live items across Mail, Documents, Workflows, and Processes
with distinct color legends, type options, and package filters.
=============================================================================
"""

import os
import io
from datetime import datetime, timedelta
import duckdb
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from pathlib import Path as _Path
_SRV_ROOT = _Path(__file__).resolve().parent.parent.parent
_SRV_CLOUD_DB = _SRV_ROOT / "aconex_cloud.duckdb"
_SRV_LOCAL_DB = _SRV_ROOT / "aconex.duckdb"

if _SRV_LOCAL_DB.exists():
    DB_PATH = str(_SRV_LOCAL_DB).replace("\\", "/")
elif _SRV_CLOUD_DB.exists():
    DB_PATH = str(_SRV_CLOUD_DB).replace("\\", "/")
else:
    DB_PATH = str(_SRV_LOCAL_DB).replace("\\", "/")


def get_db():
    return duckdb.connect(DB_PATH, read_only=True)


def get_timeframe_dates(con, timeframe: str = "today_yesterday", from_date: str = None, to_date: str = None) -> tuple[str, str, str]:
    """Returns (start_date, end_date, max_date_str) based on presets or custom dates."""
    max_dt_row = con.execute("""
        SELECT MAX(latest_dt) FROM (
            SELECT MAX(TRY_CAST(sent_date AS DATE)) as latest_dt FROM fact_mail WHERE TRY_CAST(sent_date AS DATE) <= CURRENT_DATE
            UNION ALL
            SELECT MAX(TRY_CAST(COALESCE(date_modified, revision_date) AS DATE)) as latest_dt 
            FROM dim_documents 
            WHERE TRY_CAST(COALESCE(date_modified, revision_date) AS DATE) <= CURRENT_DATE
        )
    """).fetchone()

    base_date = max_dt_row[0] if (max_dt_row and max_dt_row[0]) else datetime.now().date()
    end_date_str = base_date.strftime("%Y-%m-%d")

    # If explicit custom from/to dates are provided
    if from_date or to_date or timeframe == "custom":
        s_date = from_date.strip() if from_date and from_date.strip() else "1990-01-01"
        e_date = to_date.strip() if to_date and to_date.strip() else end_date_str
        return s_date, e_date, end_date_str

    if timeframe == "today_yesterday":
        start_date = base_date - timedelta(days=1)
    elif timeframe == "7d":
        start_date = base_date - timedelta(days=6)
    elif timeframe == "14d":
        start_date = base_date - timedelta(days=13)
    elif timeframe == "30d":
        start_date = base_date - timedelta(days=29)
    elif timeframe == "60d":
        start_date = base_date - timedelta(days=59)
    elif timeframe == "all":
        return "1990-01-01", "2099-12-31", end_date_str
    else:
        start_date = base_date - timedelta(days=1)

    return start_date.strftime("%Y-%m-%d"), end_date_str, end_date_str


def get_feed_summary(timeframe: str = "today_yesterday", from_date: str = None, to_date: str = None) -> dict:
    """Returns dynamic pills, source counts, and active types for the selected timeframe or custom range."""
    con = get_db()
    try:
        start_date, end_date, max_date = get_timeframe_dates(con, timeframe, from_date, to_date)

        # 1. Mail Active Types (with smart submittal detection for Shop Drawings & Materials)
        mail_types = con.execute(f"""
            SELECT 
                CASE 
                    WHEN mail_type IN ('Workflow Transmittal', 'Transmittal') AND (subject ILIKE '%shop%draw%' OR subject ILIKE '%-SD-%' OR subject ILIKE '%DWG%') THEN 'Shop Drawing (Transmittal)'
                    WHEN mail_type IN ('Workflow Transmittal', 'Transmittal') AND (subject ILIKE '%material%sub%' OR subject ILIKE '%material%app%' OR subject ILIKE '%-MS-%' OR subject ILIKE '%MAR-%') THEN 'Material Submittal (Transmittal)'
                    WHEN mail_type IN ('Workflow Transmittal', 'Transmittal') AND (subject ILIKE '%inspection%' OR subject ILIKE '%-IR-%' OR subject ILIKE '%WIR%') THEN 'Inspection Request (WIR/IR)'
                    WHEN mail_type IN ('Workflow Transmittal', 'Transmittal') AND (subject ILIKE '%method%statement%' OR subject ILIKE '%-MOS-%') THEN 'Method Statement (Transmittal)'
                    ELSE COALESCE(NULLIF(trim(mail_type), ''), 'General Correspondence')
                END AS m_type,
                COUNT(*) AS cnt
            FROM fact_mail
            WHERE CAST(sent_date AS DATE) >= '{start_date}' AND CAST(sent_date AS DATE) <= '{end_date}'
            GROUP BY m_type
            ORDER BY cnt DESC
        """).fetchall()

        # 2. Document Active Types
        doc_types = con.execute(f"""
            SELECT 
                COALESCE(NULLIF(trim(doc_type), ''), 'Document') AS d_type,
                COUNT(*) AS cnt
            FROM dim_documents
            WHERE CAST(COALESCE(date_modified, revision_date) AS DATE) >= '{start_date}' 
              AND CAST(COALESCE(date_modified, revision_date) AS DATE) <= '{end_date}'
            GROUP BY d_type
            ORDER BY cnt DESC
        """).fetchall()

        # 3. Workflow Active Types
        wf_types = con.execute(f"""
            SELECT 
                COALESCE(NULLIF(trim(template), ''), 'Workflow Review') AS w_type,
                COUNT(*) AS cnt
            FROM fact_workflow_steps
            WHERE CAST(COALESCE(date_completed, date_in) AS DATE) >= '{start_date}' 
              AND CAST(COALESCE(date_completed, date_in) AS DATE) <= '{end_date}'
            GROUP BY w_type
            ORDER BY cnt DESC
        """).fetchall()

        # 4. Ball-in-Court / Process Active Types
        bic_types = con.execute(f"""
            SELECT 
                COALESCE(NULLIF(trim(file_type_name), ''), 'Process Item') AS b_type,
                COUNT(*) AS cnt
            FROM fact_ball_in_court_processes
            WHERE TRY_CAST(initiated_on AS DATE) >= '{start_date}' 
              AND TRY_CAST(initiated_on AS DATE) <= '{end_date}'
            GROUP BY b_type
            ORDER BY cnt DESC
        """).fetchall()

        # Build list of active types with color and source tags
        active_types = []
        source_counts = {"mail": 0, "doc": 0, "workflow": 0, "process": 0}

        for t_name, cnt in mail_types:
            source_counts["mail"] += cnt
            active_types.append({
                "type": t_name,
                "label": f"{t_name} (Mail)",
                "source": "mail",
                "source_label": "(Mail)",
                "source_name": "Mail",
                "color": "blue",
                "count": cnt
            })

        for t_name, cnt in doc_types:
            source_counts["doc"] += cnt
            active_types.append({
                "type": t_name,
                "label": f"{t_name} (Doc)",
                "source": "doc",
                "source_label": "(Doc)",
                "source_name": "Document",
                "color": "emerald",
                "count": cnt
            })

        for t_name, cnt in wf_types:
            source_counts["workflow"] += cnt
            active_types.append({
                "type": t_name,
                "label": f"{t_name} (Workflow)",
                "source": "workflow",
                "source_label": "(Workflow)",
                "source_name": "Workflow",
                "color": "purple",
                "count": cnt
            })

        for t_name, cnt in bic_types:
            source_counts["process"] += cnt
            active_types.append({
                "type": t_name,
                "label": f"{t_name} (Process)",
                "source": "process",
                "source_label": "(Process)",
                "source_name": "Process",
                "color": "amber",
                "count": cnt
            })

        # Sort active types by count descending
        active_types.sort(key=lambda x: x["count"], reverse=True)

        # 5. Active Packages in Timeframe
        packages = con.execute(f"""
            WITH all_pkgs AS (
                SELECT project_package AS pkg FROM fact_mail WHERE CAST(sent_date AS DATE) >= '{start_date}' AND CAST(sent_date AS DATE) <= '{end_date}'
                UNION ALL
                SELECT package_name AS pkg FROM dim_documents WHERE CAST(COALESCE(date_modified, revision_date) AS DATE) >= '{start_date}' AND CAST(COALESCE(date_modified, revision_date) AS DATE) <= '{end_date}'
                UNION ALL
                SELECT project_package AS pkg FROM fact_ball_in_court_processes WHERE TRY_CAST(initiated_on AS DATE) >= '{start_date}' AND TRY_CAST(initiated_on AS DATE) <= '{end_date}'
            )
            SELECT COALESCE(NULLIF(trim(pkg), ''), 'General') AS project_id, COUNT(*) AS cnt
            FROM all_pkgs
            GROUP BY project_id
            ORDER BY cnt DESC
        """).fetchall()

        active_packages = [{"project_id": r[0], "count": r[1]} for r in packages if r[0]]

        total_all = sum(source_counts.values())

        return {
            "timeframe": timeframe,
            "start_date": start_date,
            "end_date": end_date,
            "max_date": max_date,
            "total_items": total_all,
            "source_counts": source_counts,
            "active_types": active_types,
            "active_packages": active_packages
        }
    finally:
        con.close()


def get_feed_items(
    timeframe: str = "today_yesterday",
    source: str = "all",
    item_type: str = "all",
    project: str = "all",
    search: str = None,
    from_date: str = None,
    to_date: str = None,
    discipline: str = None,
    status: str = None,
    page: int = 1,
    limit: int = 50
) -> dict:
    """Returns paginated unified feed items matching all active filters."""
    con = get_db()
    try:
        start_date, end_date, max_date = get_timeframe_dates(con, timeframe, from_date, to_date)

        queries = []

        # 1. Mail Query
        if source in ["all", "mail"]:
            queries.append(f"""
                SELECT 
                    'mail' AS source,
                    '(Mail)' AS source_label,
                    'blue' AS source_color,
                    m.mail_no AS item_id,
                    CASE 
                        WHEN m.mail_type IN ('Workflow Transmittal', 'Transmittal') AND (m.subject ILIKE '%shop%draw%' OR m.subject ILIKE '%-SD-%' OR m.subject ILIKE '%DWG%') THEN 'Shop Drawing (Transmittal)'
                        WHEN m.mail_type IN ('Workflow Transmittal', 'Transmittal') AND (m.subject ILIKE '%material%sub%' OR m.subject ILIKE '%material%app%' OR m.subject ILIKE '%-MS-%' OR m.subject ILIKE '%MAR-%') THEN 'Material Submittal (Transmittal)'
                        WHEN m.mail_type IN ('Workflow Transmittal', 'Transmittal') AND (m.subject ILIKE '%inspection%' OR m.subject ILIKE '%-IR-%' OR m.subject ILIKE '%WIR%') THEN 'Inspection Request (WIR/IR)'
                        WHEN m.mail_type IN ('Workflow Transmittal', 'Transmittal') AND (m.subject ILIKE '%method%statement%' OR m.subject ILIKE '%-MOS-%') THEN 'Method Statement (Transmittal)'
                        ELSE COALESCE(NULLIF(trim(m.mail_type), ''), 'General Correspondence')
                    END AS item_type,
                    COALESCE(m.subject, 'No Subject') AS title,
                    COALESCE(NULLIF(trim(m.project_package), ''), 'General') AS project_id,
                    COALESCE(m.from_organization, 'Unknown Org') AS organization,
                    COALESCE(m.from_name, '') AS author,
                    COALESCE(m.status, 'Active') AS status,
                    m.sent_date AS item_date,
                    CASE 
                        WHEN m.aconex_url LIKE '%MiniURL%' THEN m.aconex_url
                        WHEN l.aconex_url LIKE '%MiniURL%' THEN l.aconex_url
                        ELSE '/api/aconex/mail?mail_no=' || trim(replace(replace(m.mail_no, '"', ''), '''', ''))
                    END AS aconex_url
                FROM fact_mail m
                LEFT JOIN mail_aconex_links l ON m.mail_no = l.mail_no
                WHERE CAST(m.sent_date AS DATE) >= '{start_date}' AND CAST(m.sent_date AS DATE) <= '{end_date}'
            """)

        # 2. Document Query
        if source in ["all", "doc", "document"]:
            queries.append(f"""
                SELECT 
                    'doc' AS source,
                    '(Doc)' AS source_label,
                    'emerald' AS source_color,
                    document_no AS item_id,
                    COALESCE(NULLIF(trim(doc_type), ''), 'Document') AS item_type,
                    COALESCE(title, 'Untitled Document') AS title,
                    COALESCE(NULLIF(trim(package_name), ''), 'General') AS project_id,
                    COALESCE(review_status, 'Registered') AS organization,
                    COALESCE(created_by, '') AS author,
                    COALESCE(status, 'Current') AS status,
                    COALESCE(date_modified, revision_date) AS item_date,
                    COALESCE(NULLIF(trim(aconex_url), ''), '/api/aconex/doc?doc_no=' || TRIM(BOTH '''' FROM TRIM(BOTH '"' FROM TRIM(document_no))) || '%22') AS aconex_url
                FROM dim_documents
                WHERE CAST(COALESCE(date_modified, revision_date) AS DATE) >= '{start_date}' 
                  AND CAST(COALESCE(date_modified, revision_date) AS DATE) <= '{end_date}'
            """)

        # 3. Workflow Query
        if source in ["all", "workflow"]:
            queries.append(f"""
                SELECT 
                    'workflow' AS source,
                    '(Workflow)' AS source_label,
                    'purple' AS source_color,
                    workflow_no AS item_id,
                    COALESCE(NULLIF(trim(template), ''), 'Workflow Review') AS item_type,
                    document_title || ' - ' || step_name AS title,
                    'General' AS project_id,
                    COALESCE(assigned_to_organization, 'Hill International') AS organization,
                    COALESCE(assigned_to_name, '') AS author,
                    COALESCE(step_status, 'In Progress') AS status,
                    COALESCE(date_completed, date_in) AS item_date,
                    CASE 
                        WHEN document_no IS NOT NULL AND trim(document_no) != '' 
                        THEN '/api/aconex/doc?doc_no=' || trim(replace(replace(document_no, '"', ''), '''', ''))
                        ELSE 'https://mea.aconex.com/hub/index.html?mainTarget=%2FDocSearch%3FPROJECT_ID%3D1744833878'
                    END AS aconex_url
                FROM fact_workflow_steps
                WHERE CAST(COALESCE(date_completed, date_in) AS DATE) >= '{start_date}' 
                  AND CAST(COALESCE(date_completed, date_in) AS DATE) <= '{end_date}'
            """)

        # 4. Ball in Court Query
        if source in ["all", "process"]:
            queries.append(f"""
                SELECT 
                    'process' AS source,
                    '(Process)' AS source_label,
                    'amber' AS source_color,
                    process_ref AS item_id,
                    COALESCE(NULLIF(trim(file_type_name), ''), 'Process Item') AS item_type,
                    COALESCE(latest_activity, process_ref) AS title,
                    COALESCE(NULLIF(trim(project_package), ''), 'General') AS project_id,
                    COALESCE(contractor, 'Contractor') AS organization,
                    COALESCE(ball_in_court, '') AS author,
                    COALESCE(process_status, 'Active') AS status,
                    TRY_CAST(initiated_on AS TIMESTAMP) AS item_date,
                    aconex_url
                FROM fact_ball_in_court_processes
                WHERE TRY_CAST(initiated_on AS DATE) >= '{start_date}' 
                  AND TRY_CAST(initiated_on AS DATE) <= '{end_date}'
            """)

        if not queries:
            return {"total": 0, "page": page, "limit": limit, "items": []}

        union_sql = " UNION ALL ".join(queries)

        outer_filters = ["1=1"]
        params = []

        if item_type and item_type != "all":
            outer_filters.append("item_type = ?")
            params.append(item_type)

        if project and project != "all":
            outer_filters.append("(project_id = ? OR project_id ILIKE ?)")
            params.extend([project, f"%{project}%"])

        if search and search.strip():
            s = f"%{search.strip().lower()}%"
            outer_filters.append("(lower(item_id) LIKE ? OR lower(title) LIKE ? OR lower(organization) LIKE ? OR lower(author) LIKE ?)")
            params.extend([s, s, s, s])

        where_clause = " AND ".join(outer_filters)

        count_sql = f"""
            SELECT COUNT(*) 
            FROM ({union_sql}) u
            WHERE {where_clause}
        """
        total = con.execute(count_sql, params).fetchone()[0]

        offset = max(0, (page - 1) * limit)
        data_sql = f"""
            SELECT 
                source, source_label, source_color,
                item_id, item_type, title, project_id,
                organization, author, status,
                item_date::VARCHAR AS item_date_str,
                aconex_url
            FROM ({union_sql}) u
            WHERE {where_clause}
            ORDER BY item_date DESC NULLS LAST
            LIMIT {limit} OFFSET {offset}
        """
        rows = con.execute(data_sql, params).fetchall()

        items = []
        for r in rows:
            dt_raw = r[10]
            dt_display = dt_raw[:10] if dt_raw else "N/A"
            relative_tag = "Today" if dt_display == max_date else ("Yesterday" if dt_display == (datetime.strptime(max_date, "%Y-%m-%d").date() - timedelta(days=1)).strftime("%Y-%m-%d") else dt_display)

            items.append({
                "source": r[0],
                "source_label": r[1],
                "source_color": r[2],
                "item_id": r[3],
                "item_type": r[4],
                "title": r[5],
                "project_id": r[6],
                "organization": r[7],
                "author": r[8],
                "status": r[9],
                "item_date": dt_display,
                "relative_tag": relative_tag,
                "aconex_url": r[11]
            })

        return {
            "total": total,
            "page": page,
            "limit": limit,
            "start_date": start_date,
            "end_date": end_date,
            "items": items
        }
    finally:
        con.close()


def generate_feed_excel(
    timeframe: str = "today_yesterday",
    source: str = "all",
    item_type: str = "all",
    project: str = "all",
    search: str = None
) -> tuple[bytes, str]:
    """Generates a high-quality Excel export of the filtered daily feed with live hyperlinks."""
    res = get_feed_items(timeframe, source, item_type, project, search, page=1, limit=10000)
    items = res.get("items", [])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Daily Activity Feed"
    ws.views.sheetView[0].showGridLines = True

    # Styling constants
    font_title = Font(name="Calibri", size=15, bold=True, color="1E293B")
    font_sub = Font(name="Calibri", size=10, italic=True, color="64748B")
    font_hdr = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_body = Font(name="Calibri", size=10, color="0F172A")
    font_link = Font(name="Calibri", size=10, color="0284C7", underline="single")
    
    fill_hdr = PatternFill("solid", fgColor="1E293B")
    fill_even = PatternFill("solid", fgColor="F8FAFC")
    fill_odd = PatternFill("solid", fgColor="FFFFFF")

    border_thin = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    # Title Block
    ws.append(["Ras El Hekma Project - Daily Activity & Deliverables Feed"])
    ws.cell(1, 1).font = font_title
    
    date_info = f"Report Date Range: {res.get('start_date')} to {res.get('end_date')} | Total Items: {len(items):,} | Source Filter: {source.upper()} | Type Filter: {item_type} | Package: {project}"
    ws.append([date_info])
    ws.cell(2, 1).font = font_sub
    ws.append([])

    # Table Headers
    headers = [
        "Date", "Timeline", "Source", "Type", "Reference / Number",
        "Title / Subject", "Package", "Organization", "Author", "Status", "Aconex Hyperlink"
    ]
    ws.append(headers)
    hdr_row = ws.max_row

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(hdr_row, col_idx)
        cell.font = font_hdr
        cell.fill = fill_hdr
        cell.alignment = Alignment(horizontal="center" if col_idx in [1, 2, 3, 7, 10] else "left", vertical="center")
        cell.border = border_thin

    ws.row_dimensions[hdr_row].height = 26

    # Data Rows
    for i, it in enumerate(items):
        row_idx = ws.max_row + 1
        fill_curr = fill_even if i % 2 == 0 else fill_odd

        row_vals = [
            it.get("item_date"),
            it.get("relative_tag"),
            f"{it.get('source', '').upper()} {it.get('source_label', '')}",
            it.get("item_type"),
            it.get("item_id"),
            it.get("title"),
            it.get("project_id"),
            it.get("organization"),
            it.get("author"),
            it.get("status"),
            it.get("aconex_url")
        ]
        ws.append(row_vals)

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row_idx, col_idx)
            cell.fill = fill_curr
            cell.border = border_thin
            cell.font = font_body
            cell.alignment = Alignment(
                horizontal="center" if col_idx in [1, 2, 3, 7, 10] else "left",
                vertical="center"
            )

            # Hyperlink column
            if col_idx == 11 and it.get("aconex_url"):
                cell.font = font_link
                target_url = it["aconex_url"]
                cell.hyperlink = target_url

        ws.row_dimensions[row_idx].height = 20

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 60)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    filename = f"Aconex_Daily_Feed_{res.get('start_date')}_to_{res.get('end_date')}.xlsx"
    return out.getvalue(), filename
