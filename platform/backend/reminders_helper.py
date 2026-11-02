import io
import os
import re
import subprocess
import tempfile
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.application import MIMEApplication
from datetime import datetime

try:
    from platform.backend.package_resolver import group_dossier_by_package, DEFAULT_PMC_CC, resolve_project_package
except ImportError:
    from package_resolver import group_dossier_by_package, DEFAULT_PMC_CC, resolve_project_package


def generate_company_excel(org_name: str, dossier_data: dict) -> bytes:
    """
    Generates an executive-styled Excel workbook containing overdue items
    for the specified company across Workflows, Mails, and Lookahead,
    fully grouped by Project / Design Package with live native Aconex hyperlinks.
    """
    wb = openpyxl.Workbook()
    
    # Styles
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_title = Font(name="Calibri", size=14, bold=True, color="1F497D")
    font_data = Font(name="Calibri", size=10)
    font_bold = Font(name="Calibri", size=10, bold=True)
    font_red = Font(name="Calibri", size=10, bold=True, color="C00000")
    font_link = Font(name="Calibri", size=10, color="004B87", underline="single")
    
    fill_header_navy = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    fill_header_red = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
    fill_header_amber = PatternFill(start_color="ED7D31", end_color="ED7D31", fill_type="solid")
    fill_header_slate = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    fill_zebra = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    grouped = group_dossier_by_package(dossier_data)
    packages_data = grouped.get("packages", {})

    # -------------------------------------------------------------
    # Sheet 1: Project Package Breakdown (Summary Sheet)
    # -------------------------------------------------------------
    ws_pkg = wb.active
    ws_pkg.title = "Package Breakdown"
    
    ws_pkg.cell(1, 1, f"RAS EL HEKMA PROJECT — PACKAGE OVERDUE SUMMARY: {org_name}").font = font_title
    ws_pkg.cell(2, 1, f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')} | Total Project Packages: {len(packages_data)}").font = font_bold
    ws_pkg.row_dimensions[1].height = 26
    ws_pkg.row_dimensions[4].height = 24

    pkg_headers = [
        "Project / Design Package",
        "Overdue Review Steps",
        "Overdue Mails & RFIs",
        "Total Overdue Items",
        "Max Delay (Days)",
        "Due This Week (7d Lookahead)"
    ]
    for col_idx, h in enumerate(pkg_headers, start=1):
        cell = ws_pkg.cell(4, col_idx, h)
        cell.font = font_header
        cell.fill = fill_header_slate
        cell.alignment = align_center
        cell.border = thin_border

    r_idx = 5
    for pkg_name, p_info in packages_data.items():
        ws_pkg.row_dimensions[r_idx].height = 20
        wf_cnt = len(p_info.get("wf_steps", []))
        mail_cnt = len(p_info.get("mails", [])) + len(p_info.get("rfis", []))
        total_cnt = wf_cnt + mail_cnt
        max_d = p_info.get("max_delay", 0)
        look_cnt = len(p_info.get("lookahead", []))

        row_cells = [
            ws_pkg.cell(r_idx, 1, pkg_name),
            ws_pkg.cell(r_idx, 2, wf_cnt),
            ws_pkg.cell(r_idx, 3, mail_cnt),
            ws_pkg.cell(r_idx, 4, total_cnt),
            ws_pkg.cell(r_idx, 5, f"+{max_d} d" if max_d > 0 else "0 d"),
            ws_pkg.cell(r_idx, 6, look_cnt)
        ]
        is_even = (r_idx % 2 == 0)
        for c_i, c in enumerate(row_cells, start=1):
            c.font = font_bold if c_i in [1, 4] else font_data
            c.border = thin_border
            if is_even:
                c.fill = fill_zebra
            if c_i == 1:
                c.alignment = align_left
            elif c_i in [2, 3, 4, 5, 6]:
                c.alignment = align_center
            if c_i == 4 and total_cnt > 0:
                c.font = font_red
        r_idx += 1

    for col in ws_pkg.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_pkg.column_dimensions[col_letter].width = min(max(max_len + 4, 15), 65)

    # -------------------------------------------------------------
    # Sheet 2: Overdue Workflows (Grouped by Package)
    # -------------------------------------------------------------
    ws_wf = wb.create_sheet("Overdue Workflows")
    ws_wf.cell(1, 1, f"RAS EL HEKMA PROJECT — OVERDUE WORKFLOWS: {org_name}").font = font_title
    ws_wf.cell(2, 1, f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')} | Total Overdue Steps: {len(dossier_data.get('wf_steps', []))}").font = font_bold
    ws_wf.row_dimensions[1].height = 26
    ws_wf.row_dimensions[4].height = 24
    
    wf_headers = [
        "Project / Package",
        "Workflow No",
        "Step Name",
        "Document No",
        "Document Title",
        "Due Date",
        "Days Late",
        "Assignee",
        "Aconex Direct Link"
    ]
    for col_idx, h in enumerate(wf_headers, start=1):
        cell = ws_wf.cell(4, col_idx, h)
        cell.font = font_header
        cell.fill = fill_header_navy
        cell.alignment = align_center
        cell.border = thin_border
        
    wf_items = []
    for pkg_name, p_info in packages_data.items():
        for s in p_info.get("wf_steps", []):
            wf_items.append((pkg_name, s))
            
    r_idx = 5
    for pkg_name, s in wf_items:
        ws_wf.row_dimensions[r_idx].height = 20
        link_val = "Open in Aconex ↗" if s.get("aconex_url") else "-"
        row_cells = [
            ws_wf.cell(r_idx, 1, pkg_name),
            ws_wf.cell(r_idx, 2, s.get("workflow_no")),
            ws_wf.cell(r_idx, 3, s.get("step_name")),
            ws_wf.cell(r_idx, 4, s.get("document_no")),
            ws_wf.cell(r_idx, 5, s.get("title")),
            ws_wf.cell(r_idx, 6, s.get("due_date")),
            ws_wf.cell(r_idx, 7, s.get("days_late")),
            ws_wf.cell(r_idx, 8, s.get("user") or "Unassigned"),
            ws_wf.cell(r_idx, 9, link_val)
        ]
        is_even = (r_idx % 2 == 0)
        for c_i, c in enumerate(row_cells, start=1):
            c.font = font_data
            c.border = thin_border
            if is_even:
                c.fill = fill_zebra
            if c_i in [2, 4, 6]:
                c.alignment = align_center
            elif c_i == 7:
                c.alignment = align_center
                c.font = font_red
            elif c_i == 9 and s.get("aconex_url"):
                c.font = font_link
                c.hyperlink = s.get("aconex_url")
                c.alignment = align_center
        r_idx += 1
                
    for col in ws_wf.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_wf.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 65)

    # -------------------------------------------------------------
    # Sheet 3: Overdue Letters & RFIs (Grouped by Package)
    # -------------------------------------------------------------
    ws_mail = wb.create_sheet("Overdue Letters & RFIs")
    ws_mail.cell(1, 1, f"RAS EL HEKMA PROJECT — OVERDUE CORRESPONDENCE & RFIs: {org_name}").font = font_title
    ws_mail.cell(2, 1, f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')} | Total Mails & RFIs: {len(dossier_data.get('mails', [])) + len(dossier_data.get('rfis', []))}").font = font_bold
    ws_mail.row_dimensions[1].height = 26
    ws_mail.row_dimensions[4].height = 24
    
    mail_headers = [
        "Project / Package",
        "Mail / RFI No",
        "Type",
        "Subject",
        "From Organization",
        "Originator",
        "Sent Date",
        "Due Date",
        "Days Late",
        "Aconex Direct Link"
    ]
    for col_idx, h in enumerate(mail_headers, start=1):
        cell = ws_mail.cell(4, col_idx, h)
        cell.font = font_header
        cell.fill = fill_header_red
        cell.alignment = align_center
        cell.border = thin_border
        
    mail_items = []
    for pkg_name, p_info in packages_data.items():
        for m in p_info.get("mails", []):
            mail_items.append((pkg_name, m))
        for r in p_info.get("rfis", []):
            rfi_dict = {
                "mail_no": r.get("mail_no"), "type": "RFI", "subject": r.get("subject"),
                "from_org": r.get("from_org"), "from_name": r.get("from_name"),
                "sent_date": r.get("sent_date"), "due_date": r.get("due_date"), "days_late": r.get("days_late"),
                "aconex_url": r.get("aconex_url")
            }
            mail_items.append((pkg_name, rfi_dict))

    r_idx = 5
    for pkg_name, m in mail_items:
        ws_mail.row_dimensions[r_idx].height = 20
        link_val = "Open in Aconex ↗" if m.get("aconex_url") else "-"
        row_cells = [
            ws_mail.cell(r_idx, 1, pkg_name),
            ws_mail.cell(r_idx, 2, m.get("mail_no")),
            ws_mail.cell(r_idx, 3, m.get("type")),
            ws_mail.cell(r_idx, 4, m.get("subject")),
            ws_mail.cell(r_idx, 5, m.get("from_org")),
            ws_mail.cell(r_idx, 6, m.get("from_name")),
            ws_mail.cell(r_idx, 7, m.get("sent_date")),
            ws_mail.cell(r_idx, 8, m.get("due_date")),
            ws_mail.cell(r_idx, 9, m.get("days_late")),
            ws_mail.cell(r_idx, 10, link_val)
        ]
        is_even = (r_idx % 2 == 0)
        for c_i, c in enumerate(row_cells, start=1):
            c.font = font_data
            c.border = thin_border
            if is_even:
                c.fill = fill_zebra
            if c_i in [2, 3, 7, 8]:
                c.alignment = align_center
            elif c_i == 9:
                c.alignment = align_center
                c.font = font_red
            elif c_i == 10 and m.get("aconex_url"):
                c.font = font_link
                c.hyperlink = m.get("aconex_url")
                c.alignment = align_center
        r_idx += 1
                
    for col in ws_mail.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_mail.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 65)

    # -------------------------------------------------------------
    # Sheet 4: Lookahead 7 Days (Grouped by Package)
    # -------------------------------------------------------------
    ws_look = wb.create_sheet("Lookahead 7 Days")
    ws_look.cell(1, 1, f"RAS EL HEKMA PROJECT — TASKS DUE THIS WEEK (7-DAY LOOKAHEAD): {org_name}").font = font_title
    ws_look.cell(2, 1, f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')} | Total Tasks Due: {len(dossier_data.get('lookahead_7d', []))}").font = font_bold
    ws_look.row_dimensions[1].height = 26
    ws_look.row_dimensions[4].height = 24
    
    look_headers = [
        "Project / Package",
        "Workflow No",
        "Step Name",
        "Document No",
        "Document Title",
        "Due Date",
        "Days Remaining",
        "Assignee",
        "Aconex Direct Link"
    ]
    for col_idx, h in enumerate(look_headers, start=1):
        cell = ws_look.cell(4, col_idx, h)
        cell.font = font_header
        cell.fill = fill_header_amber
        cell.alignment = align_center
        cell.border = thin_border
        
    look_items = []
    for pkg_name, p_info in packages_data.items():
        for l in p_info.get("lookahead", []):
            look_items.append((pkg_name, l))

    r_idx = 5
    for pkg_name, l in look_items:
        ws_look.row_dimensions[r_idx].height = 20
        link_val = "Open in Aconex ↗" if l.get("aconex_url") else "-"
        row_cells = [
            ws_look.cell(r_idx, 1, pkg_name),
            ws_look.cell(r_idx, 2, l.get("workflow_no")),
            ws_look.cell(r_idx, 3, l.get("step_name")),
            ws_look.cell(r_idx, 4, l.get("document_no")),
            ws_look.cell(r_idx, 5, l.get("title")),
            ws_look.cell(r_idx, 6, l.get("due_date")),
            ws_look.cell(r_idx, 7, l.get("days_remaining")),
            ws_look.cell(r_idx, 8, l.get("user") or "Unassigned"),
            ws_look.cell(r_idx, 9, link_val)
        ]
        is_even = (r_idx % 2 == 0)
        for c_i, c in enumerate(row_cells, start=1):
            c.font = font_data
            c.border = thin_border
            if is_even:
                c.fill = fill_zebra
            if c_i in [2, 4, 6, 7]:
                c.alignment = align_center
            elif c_i == 9 and l.get("aconex_url"):
                c.font = font_link
                c.hyperlink = l.get("aconex_url")
                c.alignment = align_center
        r_idx += 1
                
    for col in ws_look.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_look.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 65)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def generate_company_html_table(target_name: str, dossier_data: dict, is_user: bool = False, org_name: str = None) -> str:
    """
    Generates a high-quality responsive HTML email template featuring deliverables
    strictly organized by Project / Design Package, styled tables with direct clickable
    Aconex hyperlinks, and the contractual action statement situated at the bottom.
    Supports both Organization-level dossiers and Individual Reviewer / Engineer reminders.
    """
    grouped = group_dossier_by_package(dossier_data)
    packages_data = grouped.get("packages", {})
    
    total_wf = len(dossier_data.get("wf_steps", []))
    total_mails = len(dossier_data.get("mails", []))
    total_rfis = len(dossier_data.get("rfis", []))
    total_look = len(dossier_data.get("lookahead_7d", []))
    total_overdue = total_wf + total_mails + total_rfis
    max_delay = dossier_data.get("max_delay_days", 0)

    header_title = "Official Technical Review Overdue Notice" if is_user else "Official PMC Weekly Overdue Notice"
    target_label = "Target Reviewer" if is_user else "Target Organization"
    target_sub = f"{target_name} ({org_name})" if (is_user and org_name) else target_name

    if is_user:
        greeting_html = f"""Dear <strong>{target_name}</strong>,<br><br>
        According to Oracle Aconex project records, the following technical review deliverables assigned to you are currently exceeding contractual turnaround times. 
        All deliverables are itemized below by <strong>Project &amp; Design Package</strong> with direct 1-click Aconex links. Immediate review and action are requested before the end of this working week."""
    else:
        greeting_html = f"""Dear <strong>{target_name}</strong> Management &amp; Controls Team,<br><br>
        This is the official PMC weekly notification regarding submittals, workflows, and correspondence exceeding contractual turnaround times in <strong>Oracle Aconex</strong>. 
        All deliverables are itemized below by <strong>Project &amp; Design Package</strong>. Immediate resolution is requested."""

    html = f"""<!DOCTYPE html>
<html>
<head>
    <meta charset="utf-8">
    <style>
        body.email-body {{ font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, Helvetica, Arial, sans-serif; margin: 0; padding: 20px; background-color: #f8fafc; color: #1e293b; line-height: 1.5; }}
        .email-container {{ max-width: 980px; margin: 0 auto; background: #ffffff; border-radius: 12px; border: 1px solid #e2e8f0; overflow: hidden; box-shadow: 0 4px 6px -1px rgba(0,0,0,0.05); }}
        .email-container .header {{ background: linear-gradient(135deg, #0f172a 0%, #1e293b 100%); padding: 24px 32px; color: #ffffff; }}
        .email-container .kpi-row {{ display: table; width: 100%; border-bottom: 1px solid #e2e8f0; background: #f8fafc; padding: 16px 24px; box-sizing: border-box; }}
        .email-container .kpi-col {{ display: table-cell; text-align: center; width: 25%; padding: 8px; border-right: 1px solid #e2e8f0; }}
        .email-container .kpi-col:last-child {{ border-right: none; }}
        .email-container .kpi-val {{ font-size: 22px; font-weight: 800; color: #0f172a; margin-top: 4px; }}
        .email-container .kpi-val-red {{ font-size: 22px; font-weight: 800; color: #dc2626; margin-top: 4px; }}
        .email-container .kpi-lbl {{ font-size: 11px; text-transform: uppercase; color: #64748b; font-weight: 600; letter-spacing: 0.5px; }}
        .email-container .content {{ padding: 28px 32px; }}
        .email-container .pkg-header {{ background: #0f172a; color: #f8fafc; padding: 12px 18px; border-radius: 8px; margin-top: 28px; margin-bottom: 12px; display: table; width: 100%; box-sizing: border-box; }}
        .email-container .pkg-title {{ display: table-cell; vertical-align: middle; font-size: 14px; font-weight: 800; letter-spacing: 0.3px; }}
        .email-container .pkg-badge {{ display: table-cell; vertical-align: middle; text-align: right; font-size: 11px; font-weight: 700; color: #fecaca; }}
        .email-container table {{ width: 100%; border-collapse: collapse; font-size: 12px; margin-bottom: 20px; }}
        .email-container th {{ background: #1e293b; color: #f8fafc; text-align: left; padding: 9px 12px; font-weight: 600; font-size: 11px; text-transform: uppercase; letter-spacing: 0.5px; }}
        .email-container td {{ padding: 9px 12px; border-bottom: 1px solid #e2e8f0; vertical-align: top; }}
        .email-container tr:nth-child(even) {{ background: #f8fafc; }}
        .email-container .tag-delay {{ background: #fef2f2; color: #b91c1c; font-weight: 700; padding: 3px 8px; border-radius: 4px; border: 1px solid #fecaca; font-size: 11px; display: inline-block; }}
        .email-container .tag-type {{ font-size: 11px; background: #e2e8f0; color: #334155; padding: 2px 6px; border-radius: 4px; font-weight: 600; }}
        .email-container .link-btn {{ color: #0284c7; font-weight: 600; text-decoration: none; }}
        .email-container .link-btn:hover {{ text-decoration: underline; }}
        .email-container .notice-box {{ background: #fffbeb; border: 1px solid #fde68a; border-radius: 8px; padding: 18px; font-size: 13px; color: #92400e; margin-top: 30px; line-height: 1.6; }}
        .email-container .footer {{ background: #f1f5f9; padding: 20px 32px; font-size: 11px; color: #64748b; border-top: 1px solid #e2e8f0; }}
    </style>
</head>
<body class="email-body">
    <div class="email-container">
        <div class="header">
            <div style="font-size: 12px; font-weight: 700; color: #38bdf8; letter-spacing: 1px; text-transform: uppercase; margin-bottom: 4px;">
                Ras El Hekma Project Controls &amp; Governance
            </div>
            <h1 style="margin: 0; font-size: 22px; font-weight: 800;">{header_title}</h1>
            <div style="margin-top: 8px; font-size: 14px; color: #cbd5e1;">
                {target_label}: <strong style="color: #ffffff;">{target_sub}</strong>
            </div>
        </div>

        <div class="kpi-row">
            <div class="kpi-col">
                <div class="kpi-lbl">Total Overdue</div>
                <div class="kpi-val-red">{total_overdue}</div>
            </div>
            <div class="kpi-col">
                <div class="kpi-lbl">Max Delay</div>
                <div class="kpi-val-red">+{max_delay} Days</div>
            </div>
            <div class="kpi-col">
                <div class="kpi-lbl">Review Workflows</div>
                <div class="kpi-val">{total_wf}</div>
            </div>
            <div class="kpi-col">
                <div class="kpi-lbl">Due This Week (7d)</div>
                <div class="kpi-val" style="color: #d97706;">{total_look}</div>
            </div>
        </div>

        <div class="content">
            <p style="font-size: 13px; color: #475569; margin-top: 0; margin-bottom: 16px;">
                {greeting_html}
            </p>
    """

    if not packages_data:
        html += """<div style="text-align: center; color: #16a34a; padding: 32px; font-weight: bold;">
            ✓ Congratulations: No pending overdue deliverables recorded for your organization!
        </div>"""
    else:
        for pkg_name, p_info in packages_data.items():
            wf_list = p_info.get("wf_steps", [])
            mail_list = p_info.get("mails", [])
            rfi_list = p_info.get("rfis", [])
            look_list = p_info.get("lookahead", [])
            pkg_total = len(wf_list) + len(mail_list) + len(rfi_list)
            pkg_max_d = p_info.get("max_delay", 0)

            html += f"""
            <div class="pkg-header">
                <div class="pkg-title">📦 {pkg_name}</div>
                <div class="pkg-badge">
                    {pkg_total} Overdue Item(s) &bull; Max Delay: +{pkg_max_d} d
                </div>
            </div>
            """

            # 1. Workflow steps table
            if wf_list:
                html += """
                <table>
                    <thead>
                        <tr>
                            <th style="width: 14%;">Workflow No</th>
                            <th style="width: 28%;">Review Step &amp; Document Title</th>
                            <th style="width: 24%;">Document No</th>
                            <th style="width: 14%;">Assignee</th>
                            <th style="width: 10%;">Due Date</th>
                            <th style="width: 10%; text-align: right;">Delay</th>
                        </tr>
                    </thead>
                    <tbody>
                """
                for s in wf_list[:20]:
                    doc_link = s.get("aconex_url")
                    doc_html = f'<a href="{doc_link}" class="link-btn" target="_blank" title="Open directly in Aconex">{s.get("document_no")} ↗</a>' if doc_link else f'<strong>{s.get("document_no")}</strong>'
                    html += f"""
                        <tr>
                            <td><strong style="color:#0f172a;">{s.get("workflow_no")}</strong></td>
                            <td>{s.get("step_name")}<br><span style="font-size: 11px; color:#64748b;">{(s.get("title") or "")[:50]}</span></td>
                            <td>{doc_html}</td>
                            <td>{s.get("user") or "Unassigned"}</td>
                            <td>{s.get("due_date") or "--"}</td>
                            <td style="text-align: right;"><span class="tag-delay">+{s.get("days_late")} d</span></td>
                        </tr>
                    """
                if len(wf_list) > 20:
                    html += f"""<tr><td colspan="6" style="text-align:center; font-style:italic; color:#64748b; padding:10px;">... and {len(wf_list)-20} additional overdue steps for this package in attached Excel.</td></tr>"""
                html += "</tbody></table>"

            # 2. Letters & RFIs table
            all_pkg_mails = mail_list + [{
                "mail_no": r.get("mail_no"), "type": "RFI", "subject": r.get("subject"),
                "from_org": r.get("from_org"), "due_date": r.get("due_date"), "days_late": r.get("days_late"),
                "aconex_url": r.get("aconex_url")
            } for r in rfi_list]

            if all_pkg_mails:
                html += """
                <table>
                    <thead>
                        <tr>
                            <th style="width: 18%;">Mail / RFI No</th>
                            <th style="width: 10%;">Type</th>
                            <th style="width: 38%;">Subject</th>
                            <th style="width: 14%;">From Org</th>
                            <th style="width: 10%;">Due Date</th>
                            <th style="width: 10%; text-align: right;">Delay</th>
                        </tr>
                    </thead>
                    <tbody>
                """
                for m in all_pkg_mails[:15]:
                    m_no = m.get('mail_no', '')
                    aconex_mail_url = m.get("aconex_url") or f"https://mea.aconex.com/MailSearch?searchQuery=mailNumber%3A%22{m_no}%22&PROJECT_ID=1744833878"
                    html += f"""
                        <tr>
                            <td><a href="{aconex_mail_url}" class="link-btn" target="_blank" title="Open Mail in Aconex"><strong>{m.get("mail_no")}</strong> ↗</a></td>
                            <td><span class="tag-type">{m.get("type")}</span></td>
                            <td>{(m.get("subject") or "")[:55]}</td>
                            <td>{m.get("from_org")}</td>
                            <td>{m.get("due_date") or "--"}</td>
                            <td style="text-align: right;"><span class="tag-delay">+{m.get("days_late")} d</span></td>
                        </tr>
                    """
                if len(all_pkg_mails) > 15:
                    html += f"""<tr><td colspan="6" style="text-align:center; font-style:italic; color:#64748b; padding:10px;">... and {len(all_pkg_mails)-15} additional correspondence items in attached Excel.</td></tr>"""
                html += "</tbody></table>"

            # 3. Lookahead table for this package
            if look_list:
                html += """
                <table>
                    <thead>
                        <tr>
                            <th style="width: 15%;">Lookahead No</th>
                            <th style="width: 35%;">Step &amp; Document</th>
                            <th style="width: 30%;">Document Title</th>
                            <th style="width: 10%;">Due Date</th>
                            <th style="width: 10%; text-align: right;">Remaining</th>
                        </tr>
                    </thead>
                    <tbody>
                """
                for l in look_list[:10]:
                    doc_link = l.get("aconex_url")
                    doc_html = f'<a href="{doc_link}" class="link-btn" target="_blank">{l.get("document_no")} ↗</a>' if doc_link else l.get("document_no")
                    html += f"""
                        <tr>
                            <td><strong>{l.get("workflow_no")}</strong></td>
                            <td>{l.get("step_name")}<br><span style="font-size:11px;">{doc_html}</span></td>
                            <td>{(l.get("title") or "")[:45]}</td>
                            <td>{l.get("due_date") or "--"}</td>
                            <td style="text-align: right; color:#d97706; font-weight:700;">{l.get("days_remaining")} d left</td>
                        </tr>
                    """
                html += "</tbody></table>"

    # Contractual Action Notice and Footer
    target_clean = (org_name or target_name or "Project").strip()
    safe_name = re.sub(r'[^a-zA-Z0-9_\-]', '_', target_clean)
    excel_fname = f"{safe_name}_Overdue_Tasks_{datetime.now().strftime('%Y%m%d')}.xlsx"
    html += f"""
            <div class="notice-box">
                <strong>⚠️ Contractual Action Required:</strong><br>
                Please ensure all overdue deliverables listed above are updated, commented, or completed in Oracle Aconex before the end of this working week.
                <br><br>
                <em>Note: The full comprehensive itemized register with live Aconex links and package breakdowns is attached in the accompanying Excel spreadsheet ({excel_fname}).</em>
            </div>
        </div>

        <div class="footer">
            <strong>Project Management &amp; Controls (PMC) | Hill International</strong><br>
            Ras El Hekma Urban Development Project | Modon Real Estate<br>
            Platform Access: <a href="http://localhost:8000" style="color:#0284c7;">http://localhost:8000</a>
        </div>
    </div>
</body>
</html>
"""
    return html


def generate_company_eml(org_name: str, dossier_data: dict, recipient_email: str = None, cc_email: str = None) -> bytes:
    """
    Creates a complete RFC 822 .eml message file with HTML body and attached Excel spreadsheet.
    Sets X-Unsent: 1 so Outlook opens directly in Draft / Compose mode ready to send!
    ZERO fake emails.
    """
    subject = dossier_data.get("subject", f"[PMC OVERDUE DISPATCH] Ras El Hekma Project - {org_name}")
    recipients = recipient_email or dossier_data.get("recipient_email", "")
    cc_value = cc_email.strip() if cc_email and cc_email.strip() else ""
    
    msg = MIMEMultipart("mixed")
    msg["Subject"] = subject
    msg["From"] = "PMC Project Controls <DeanEnglish@hillintl.com>"
    msg["To"] = recipients
    if cc_value:
        msg["Cc"] = cc_value
    msg["Date"] = datetime.now().strftime("%a, %d %b %Y %H:%M:%S +0300")
    # Tells Outlook to open directly in Compose/Draft mode
    msg["X-Unsent"] = "1"
    
    # Alternative body (plain + HTML)
    alt_part = MIMEMultipart("alternative")
    plain_text = dossier_data.get("body_text", "")
    is_user = dossier_data.get("is_user", False)
    actual_org = dossier_data.get("org_name", org_name)
    html_text = generate_company_html_table(org_name, dossier_data, is_user=is_user, org_name=actual_org)
    
    alt_part.attach(MIMEText(plain_text, "plain", "utf-8"))
    alt_part.attach(MIMEText(html_text, "html", "utf-8"))
    msg.attach(alt_part)
    
    # Attach Excel
    excel_bytes = generate_company_excel(org_name, dossier_data)
    safe_name = re.sub(r'[^a-zA-Z0-9_\-]', '_', org_name)
    pkg_suffix = ""
    if dossier_data.get("selected_package"):
        safe_pkg = re.sub(r'[^a-zA-Z0-9_\-]', '_', dossier_data["selected_package"])[:25]
        pkg_suffix = f"_{safe_pkg}"
    filename = f"{safe_name}{pkg_suffix}_Overdue_Tasks_{datetime.now().strftime('%Y%m%d')}.xlsx"
    part_att = MIMEApplication(excel_bytes, _subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    part_att.add_header('Content-Disposition', 'attachment', filename=filename)
    msg.attach(part_att)
    
    return msg.as_bytes()


def open_company_outlook(org_name: str, dossier_data: dict, recipient_email: str = None, cc_email: str = None) -> str:
    """
    Generates the RFC 822 .eml draft with X-Unsent: 1, saves to temporary folder,
    saves the accompanying Excel spreadsheet to Exports folder,
    and directly launches Microsoft Outlook!
    """
    eml_bytes = generate_company_eml(org_name, dossier_data, recipient_email=recipient_email, cc_email=cc_email)
    safe_name = re.sub(r'[^a-zA-Z0-9_\-]', '_', org_name)
    
    # Save a permanent copy of the Excel spreadsheet in Exports
    exports_dir = r"C:\Users\bimst\Downloads\ACONEX\Exports"
    try:
        os.makedirs(exports_dir, exist_ok=True)
        excel_bytes = generate_company_excel(org_name, dossier_data)
        excel_path = os.path.join(exports_dir, f"{safe_name}_Overdue_Tasks_{datetime.now().strftime('%Y%m%d')}.xlsx")
        with open(excel_path, "wb") as f:
            f.write(excel_bytes)
    except Exception as e:
        print(f"Warning: could not save Excel to Exports: {e}")

    temp_dir = tempfile.gettempdir()
    eml_path = os.path.join(temp_dir, f"PMC_Overdue_Draft_{safe_name}.eml")
    with open(eml_path, "wb") as f:
        f.write(eml_bytes)

    # Launch Outlook via os.startfile (which executes the system's registered .eml handler)
    try:
        os.startfile(eml_path)
    except Exception as e:
        print(f"os.startfile warning: {e}")
        olk_exe = r"C:\Users\bimst\AppData\Local\Microsoft\WindowsApps\olk.exe"
        if os.path.exists(olk_exe):
            subprocess.Popen([olk_exe, eml_path])

    return eml_path
