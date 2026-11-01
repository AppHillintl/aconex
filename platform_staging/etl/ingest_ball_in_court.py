import os
import re
import datetime
import openpyxl
import pypdf
import duckdb

DB_PATH = r"C:\Users\bimst\Downloads\ACONEX\aconex.duckdb"

EXCEL_FILES = [
    {
        "file": r"C:\Users\bimst\Downloads\ACONEX\Data\0006_PROCESS DETAILS TABLE MD2.xlsx",
        "fallback": r"C:\Users\bimst\Downloads\0006_PROCESS DETAILS TABLE MD2.xlsx",
        "file_type": "RFI",
        "file_type_name": "RFI - Request For Information",
        "file_type_ar": "طلبات المعلومات (RFI)",
        "source_label": "MD2_RFI"
    },
    {
        "file": r"C:\Users\bimst\Downloads\ACONEX\Data\0006_PROCESS DETAILS TABLE MD2 (1).xlsx",
        "fallback": r"C:\Users\bimst\Downloads\0006_PROCESS DETAILS TABLE MD2 (1).xlsx",
        "file_type": "HSNCR",
        "file_type_name": "HSNCR - Health & Safety NCR",
        "file_type_ar": "تقارير عدم مطابقة السلامة (HSNCR)",
        "source_label": "MD2_HSNCR"
    },
    {
        "file": r"C:\Users\bimst\Downloads\ACONEX\Data\0006_PROCESS DETAILS TABLE MD2 (2).xlsx",
        "fallback": r"C:\Users\bimst\Downloads\0006_PROCESS DETAILS TABLE MD2 (2).xlsx",
        "file_type": "NCR",
        "file_type_name": "NCR - Quality Non-Conformance Report",
        "file_type_ar": "تقارير عدم المطابقة للجودة (NCR)",
        "source_label": "MD2_NCR"
    },
    {
        "file": r"C:\Users\bimst\Downloads\ACONEX\Data\0006_PROCESS DETAILS TABLE MD2 (3).xlsx",
        "fallback": r"C:\Users\bimst\Downloads\0006_PROCESS DETAILS TABLE MD2 (3).xlsx",
        "file_type": "HSE",
        "file_type_name": "HSE - Health & Safety Observations",
        "file_type_ar": "ملاحظات السلامة والبيئة (HSE)",
        "source_label": "MD2_HSE"
    },
    {
        "file": r"C:\Users\bimst\Downloads\ACONEX\Data\0006_PROCESS DETAILS TABLE MD2 (4).xlsx",
        "fallback": r"C:\Users\bimst\Downloads\0006_PROCESS DETAILS TABLE MD2 (4).xlsx",
        "file_type": "TRANSMITTAL",
        "file_type_name": "Transmittals - Submittals & Deliverables",
        "file_type_ar": "سجلات الإحالات والمراسلات (Transmittals)",
        "source_label": "MD2_TRANSMITTAL"
    }
]

PDF_SHAKER_PATH = r"C:\Users\bimst\Downloads\0006_PROCESS DETAILS TABLE MD2.pdf"

def classify_package(proc_ref, mail_no):
    t = f"{proc_ref or ''} {mail_no or ''}".upper()
    if any(k in t for k in ['CCCEGY', '3149', '3270', 'DP02', 'RHBW']):
        return 'DP02', 'CCC / ROWAD JV'
    if any(k in t for k in ['3271', 'RHD3', 'DP03-WEST', 'DP3-WEST']):
        return 'DP03-West', 'CCC-ROWAD JV'
    if any(k in t for k in ['INN-RED', '3199', 'DP03-EAST', 'DP3-EAST', 'INR', 'INNRED']):
        return 'DP03-East', 'Innovo-Redcon JV'
    if any(k in t for k in ['ACCE', '3204', 'DP04']):
        return 'DP04', 'Arabian Construction Company (ACC)'
    if any(k in t for k in ['-DP5-', '3205', 'DP05']):
        return 'DP05', 'Orascom Construction SAE'
    if any(k in t for k in ['HAC-RFI', '3206', 'DP06', 'REH/DP6', 'HYH']):
        return 'DP06', 'Hassan Allam Construction'
    if any(k in t for k in ['CSCEC', '3207', 'DP07']):
        return 'DP07', 'CSCEC Egypt'
    if any(k in t for k in ['OTC-JV', '3045', 'RO-RFI', 'RO PLANT']):
        return 'RO Plant', 'Orascom-Trojan-Cox JV'
    if any(k in t for k in ['EERM', '3044', 'MSS', 'MMS', 'SEWEDY']):
        return 'MMS / Substation', 'Elsewedy & Rowad Consortium'
    if any(k in t for k in ['-STP-', '3046', 'SEWAGE']):
        return 'STP', 'Hassan Allam Construction'
    if any(k in t for k in ['AR01', '3049', 'AIRPORT', 'AIR PORT', 'REHIA']):
        return 'Airport', 'Orascom Construction SAE'
    if any(k in t for k in ['P0801', '3200', 'GRAND CANAL', 'GRAND LAGOON']):
        return 'Grand Canal', 'Hassan Allam Construction'
    if any(k in t for k in ['P1201', '3107', 'WEST CANAL']):
        return 'West Canal', 'Orascom Construction SAE'
    if any(k in t for k in ['SW00', '3048', 'INFRA']):
        return 'Infra', 'HAC / HARB JV'
    if any(k in t for k in ['AED-RFI', 'AEDAS']):
        return 'Aedas Design Packages', 'Aedas Limited'
    if any(k in t for k in ['OBMI-RFI', 'OBMI', 'OBMI-GNC']):
        return 'OBMI Consultancy Packages', 'OBMI'
    if any(k in t for k in ['SSH-RFI', '3213', 'DP08', 'SOUTH GOLF', 'SDMEKU']):
        return 'DP08 - South Golf', 'Dar SSH International'
    if any(k in t for k in ['DARR-RFI', 'DAR']):
        return 'Dar Consultancy Packages', 'DAR AL HANDASAH'
    if 'CRAFTHS' in t:
        return 'Hospitality Packages', 'Craft Hospitality'
    if 'ECG-AD' in t or 'SHKECGCR' in t:
        return 'General Super-Consultancy', 'Shaker-ECG Consortium'
    return 'General Project Processes', 'PMC / Modon'

def parse_all_excel_files():
    all_records = []
    for cfg in EXCEL_FILES:
        path = cfg['file'] if os.path.exists(cfg['file']) else cfg['fallback']
        if not os.path.exists(path):
            print(f"Excel file not found at {path}")
            continue

        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.active
        file_records = 0

        for r in range(3, ws.max_row + 1):
            pref = ws.cell(r, 1).value
            mno = ws.cell(r, 11).value
            bic = ws.cell(r, 8).value
            if not pref and not mno and not bic:
                continue

            init_on = ws.cell(r, 2).value
            if isinstance(init_on, datetime.datetime):
                init_on_str = init_on.strftime('%Y-%m-%d')
            elif init_on:
                init_on_str = str(init_on)[:10]
            else:
                init_on_str = None

            proc_status = str(ws.cell(r, 3).value).strip() if ws.cell(r, 3).value else 'Outstanding'
            duration = float(ws.cell(r, 4).value) if ws.cell(r, 4).value is not None else None
            delay = float(ws.cell(r, 5).value) if ws.cell(r, 5).value is not None else None
            latest_act = str(ws.cell(r, 6).value).strip() if ws.cell(r, 6).value else cfg['file_type_name']
            mail_status = str(ws.cell(r, 7).value).strip() if ws.cell(r, 7).value else 'Outstanding'
            ball_in_court = str(bic).strip() if bic else 'Hill Inter'
            action = str(ws.cell(r, 9).value).strip() if ws.cell(r, 9).value else 'To Respond'
            overdue_days = float(ws.cell(r, 10).value) if ws.cell(r, 10).value is not None else None
            mail_no_val = str(mno).strip() if mno else (str(pref).strip() if pref else None)

            p_ref_str = str(pref).strip() if pref else mail_no_val
            pkg, contractor = classify_package(p_ref_str, mail_no_val)

            all_records.append({
                'process_ref': p_ref_str,
                'initiated_on': init_on_str,
                'process_status': proc_status,
                'duration': duration,
                'delay': delay,
                'latest_activity': latest_act,
                'mail_status': mail_status,
                'ball_in_court': ball_in_court,
                'action': action,
                'overdue_days': overdue_days,
                'mail_no': mail_no_val,
                'project_package': pkg,
                'contractor': contractor,
                'source': cfg['source_label'],
                'file_type': cfg['file_type'],
                'file_type_name': cfg['file_type_name'],
                'file_type_ar': cfg['file_type_ar'],
                'source_file': os.path.basename(path)
            })
            file_records += 1

        print(f"Parsed {file_records} records from [{cfg['file_type']}] {os.path.basename(path)}")

    return all_records

def parse_pdf_shaker():
    if not os.path.exists(PDF_SHAKER_PATH):
        print(f"PDF not found at {PDF_SHAKER_PATH}")
        return []

    reader = pypdf.PdfReader(PDF_SHAKER_PATH)
    records = []
    full_text = ""
    for page in reader.pages:
        full_text += page.extract_text() + "\n"

    normalized = re.sub(r'[ \t]+', ' ', full_text)
    lines = normalized.split('\n')
    chunks = []
    current_chunk = ""
    for line in lines:
        l = line.strip()
        if not l or 'Processes' in l or 'Process Reference' in l:
            continue
        if re.match(r'^[A-Z0-9\.\-_]{4,}\s+\d{2}-\d{2}-\d{4}', l):
            if current_chunk:
                chunks.append(current_chunk)
            current_chunk = l
        else:
            if current_chunk:
                current_chunk += " " + l

    if current_chunk:
        chunks.append(current_chunk)

    for c in chunks:
        tokens = c.split()
        if len(tokens) >= 8:
            p_ref = tokens[0]
            init_on_raw = tokens[1]
            try:
                parts = init_on_raw.split('-')
                if len(parts) == 3:
                    init_on_str = f"{parts[2]}-{parts[0]}-{parts[1]}"
                else:
                    init_on_str = init_on_raw
            except:
                init_on_str = init_on_raw

            p_status = tokens[2]
            try:
                dur = float(tokens[3])
            except:
                dur = None

            m_no = tokens[-1]
            try:
                ov_days = float(tokens[-2])
            except:
                ov_days = dur

            act = "Expecting Response"
            if "To Respond" in c:
                act = "To Respond"

            m_status = "Overdue" if "Overdue" in c else "Outstanding"
            bic = "Shaker-ECG consortium RASELHEKMA"
            if "Aedas" in c:
                bic = "Shaker-ECG consortium RASELHEKMA,Aedas Limited"
            elif "Shaker Egypt" in c:
                bic = "Shaker-ECG consortium RASELHEKMA,Shaker Egypt"

            latest = "Request For Information"
            if "Response to RFI" in c:
                latest = "Response to RFI"

            pkg, contractor = classify_package(p_ref, m_no)

            records.append({
                'process_ref': p_ref,
                'initiated_on': init_on_str,
                'process_status': p_status,
                'duration': dur,
                'delay': dur,
                'latest_activity': latest,
                'mail_status': m_status,
                'ball_in_court': bic,
                'action': act,
                'overdue_days': ov_days,
                'mail_no': m_no,
                'project_package': pkg,
                'contractor': contractor,
                'source': 'PDF_Shaker_ECG',
                'file_type': 'RFI',
                'file_type_name': 'RFI - Request For Information',
                'file_type_ar': 'طلبات المعلومات (RFI)',
                'source_file': '0006_PROCESS DETAILS TABLE MD2.pdf'
            })

    print(f"Parsed {len(records)} records from Shaker-ECG PDF.")
    return records

def ingest_to_duckdb():
    print("Ingesting ALL 5 Process Details Tables to DuckDB...")
    excel_records = parse_all_excel_files()
    pdf_records = parse_pdf_shaker()

    all_records = excel_records + pdf_records
    print(f"Total unified records to ingest across 5 types: {len(all_records)}")

    con = duckdb.connect(DB_PATH, read_only=False)

    con.execute("""
    CREATE OR REPLACE TABLE fact_ball_in_court_processes (
        process_ref VARCHAR,
        initiated_on VARCHAR,
        process_status VARCHAR,
        duration DOUBLE,
        delay DOUBLE,
        latest_activity VARCHAR,
        mail_status VARCHAR,
        ball_in_court VARCHAR,
        action VARCHAR,
        overdue_days DOUBLE,
        mail_no VARCHAR,
        project_package VARCHAR,
        contractor VARCHAR,
        is_overdue BOOLEAN,
        aconex_url VARCHAR,
        source VARCHAR,
        file_type VARCHAR,
        file_type_name VARCHAR,
        file_type_ar VARCHAR,
        source_file VARCHAR
    );
    """)

    insert_rows = []
    seen = set()
    for r in all_records:
        key = (r['file_type'], r['process_ref'], r['mail_no'], r['ball_in_court'], r['action'])
        if key in seen:
            continue
        seen.add(key)

        is_ov = (
            r['process_status'] == 'Overdue' or 
            r['mail_status'] == 'Overdue' or 
            (r['delay'] is not None and r['delay'] > 0) or 
            (r['overdue_days'] is not None and r['overdue_days'] > 0)
        )

        target_no = r['mail_no'] or r['process_ref']
        url = f"/api/aconex/mail?mail_no={target_no}"

        insert_rows.append((
            r['process_ref'],
            r['initiated_on'],
            r['process_status'],
            r['duration'],
            r['delay'],
            r['latest_activity'],
            r['mail_status'],
            r['ball_in_court'],
            r['action'],
            r['overdue_days'],
            r['mail_no'],
            r['project_package'],
            r['contractor'],
            is_ov,
            url,
            r['source'],
            r['file_type'],
            r['file_type_name'],
            r['file_type_ar'],
            r['source_file']
        ))

    con.executemany("""
    INSERT INTO fact_ball_in_court_processes VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
    """, insert_rows)

    print(f"\nInserted {len(insert_rows)} unique Ball in Court process records into DuckDB table 'fact_ball_in_court_processes'!")

    # Summary per file type
    print("\n--- Breakdown by Process / File Type in DuckDB ---")
    type_stats = con.execute("""
        SELECT 
            file_type, file_type_name, count(*) as total, 
            count(CASE WHEN is_overdue THEN 1 END) as overdue,
            count(CASE WHEN action = 'To Respond' THEN 1 END) as to_respond
        FROM fact_ball_in_court_processes
        GROUP BY file_type, file_type_name
        ORDER BY count(*) DESC
    """).fetchall()

    for row in type_stats:
        print(f"  [{row[0]}] {row[1]:40}: Total = {row[2]:4}, Overdue = {row[3]:4}, To Respond = {row[4]:4}")

    total_stat = con.execute("""
        SELECT 
            count(*) as total,
            count(CASE WHEN is_overdue THEN 1 END) as overdue,
            count(CASE WHEN ball_in_court ILIKE '%Hill%' THEN 1 END) as on_hill,
            count(CASE WHEN ball_in_court ILIKE '%Shaker%' OR ball_in_court ILIKE '%ECG%' THEN 1 END) as on_shaker_ecg
        FROM fact_ball_in_court_processes
    """).fetchone()
    print("\nGrand Total Summary:", total_stat)

    con.close()

if __name__ == '__main__':
    ingest_to_duckdb()
